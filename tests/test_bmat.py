"""Pruebas de los reportes BMAT: cálculo semanal, clasificación de tracks,
histórico y los cuatro reportes (familias A y B).

Corre con: pytest tests/test_bmat.py -v
"""
from datetime import date

import openpyxl
import pandas as pd
import pytest

from src import bmat, bmat_calculo, bmat_proceso, config, history
from src import bmat_clasificacion as clasif


@pytest.fixture(autouse=True)
def db_temporal(tmp_path, monkeypatch):
    monkeypatch.setattr(history, "DB_PATH", tmp_path / "test_universal_data.db")
    # Nada de semillas reales ni de la lista real de titularidad en las
    # pruebas: cada una arma exactamente lo que necesita.
    monkeypatch.setattr(bmat_calculo, "SEED_BMAT_CSV", tmp_path / "no_hay_semilla.csv.gz")
    monkeypatch.setattr(clasif, "SEED_CLASIFICACION_CSV", tmp_path / "no_hay_clasificacion.csv.gz")
    monkeypatch.setattr(config, "BMAT_TITULARIDAD_XLSX", tmp_path / "no_hay_titularidad.xlsx")


# --- utilidades ---------------------------------------------------------------

def _fila_hist(pais, anio, semana, banda, sello, tracks, streams, pct):
    return {"anio": anio, "semana": semana, "country_code": pais, "banda": banda,
            "label_group": sello, "tracks": tracks, "streams_millones": streams, "pct_streams": pct}


def _banda_hist(pais, anio, semana, banda, reparto):
    """reparto = {sello: (tracks, streams)}; el % se calcula."""
    total = sum(s for _, s in reparto.values())
    return [_fila_hist(pais, anio, semana, banda, sello, *reparto.get(sello, (0, 0.0)),
                       reparto.get(sello, (0, 0.0))[1] / total)
            for sello in config.BMAT_LABELS]


def _sembrar(tmp_path, filas):
    csv = tmp_path / "semilla.csv.gz"
    pd.DataFrame(filas).to_csv(csv, index=False, compression="gzip")
    return bmat_calculo.asegurar_semilla(csv)


def _fuente(filas):
    """filas = [(posicion, isrcs, distribuidora, streams, track, artista), ...]"""
    return pd.DataFrame([{
        config.BMAT_COL_POSICION: p, "Track": t, "Artista": a, config.BMAT_COL_ISRCS: i,
        config.BMAT_COL_DISTRIBUIDORA: d, config.BMAT_COL_DISQUERA: "X",
        config.BMAT_COL_STREAMS: s, config.BMAT_COL_MOVIMIENTO: "↑1",
        config.BMAT_COL_LANZAMIENTO: "2025-01-01",
    } for p, i, d, s, t, a in filas])


def _tabla_clasif(filas):
    """filas = [(isrc, disqueras, distribuidora)]"""
    return pd.DataFrame([{"isrc": i, "disqueras": q, "origen": clasif.ORIGEN_SEMILLA, "track": "",
                          "artista": "", "disquera_original": "", "distribuidora_original": d,
                          "actualizado": "semana 35 de 2026"} for i, q, d in filas])


# --- fuente ---------------------------------------------------------------------

def test_identificar_archivo():
    assert bmat_calculo.identificar_archivo("WK35-CO.xlsx") == (35, "CO")
    assert bmat_calculo.identificar_archivo("C:/x/WK36-CAM.xlsx") == (36, "CAM")
    assert bmat_calculo.identificar_archivo("WK36-PA.xlsx") == (36, "PA")
    assert bmat_calculo.identificar_archivo("wk7-pn.xlsx") == (7, "PA")  # Panamá, como en los reportes
    with pytest.raises(ValueError):
        bmat_calculo.identificar_archivo("Top 3000 BMAT CR a sem 35.xlsx")


def test_anio_de_la_semana():
    assert bmat_calculo.anio_de_la_semana(36, date(2026, 9, 21)) == 2026
    assert bmat_calculo.anio_de_la_semana(52, date(2027, 1, 8)) == 2026


def test_buscar_fuentes_rechaza_semanas_mezcladas(tmp_path):
    for nombre in ["WK36-CO.xlsx", "WK36-PE.xlsx", "notas.xlsx"]:
        (tmp_path / nombre).write_bytes(b"")
    assert set(bmat_calculo.buscar_fuentes(tmp_path, ["CO", "PE"])) == {"CO", "PE"}
    assert set(bmat_calculo.buscar_fuentes(tmp_path, ["CO"])) == {"CO"}
    (tmp_path / "WK35-PE.xlsx").write_bytes(b"")
    with pytest.raises(ValueError, match="varias semanas"):
        bmat_calculo.buscar_fuentes(tmp_path, ["CO", "PE"])


# --- cálculo --------------------------------------------------------------------

def test_calcular_bandas_por_sello_y_banda():
    df = _fuente([(p, f"I{p}", "d", 1e6 * (61 - p), "t", "a") for p in range(1, 61)])
    df["Sello"] = ["Universal" if p <= 5 else ("Sony" if p <= 20 else "Independientes") for p in range(1, 61)]
    r = bmat_calculo.calcular_bandas(df, "CAM")
    assert sorted(r["banda"].unique()) == [10, 50]      # la fuente tiene 60 filas
    top10 = r[r["banda"] == 10].set_index("label_group")
    assert top10.loc["Universal", "tracks"] == 5 and top10.loc["Sony", "tracks"] == 5
    assert top10.loc["Universal", "streams_millones"] == pytest.approx(60 + 59 + 58 + 57 + 56)
    assert top10["pct_streams"].sum() == pytest.approx(1.0)
    assert top10.loc["ADA Music", "tracks"] == 0


def test_titularidad_compartida_reparte_streams_pero_no_tracks(tmp_path):
    lista = tmp_path / "tit.xlsx"
    pd.DataFrame([{"Track": "Dakiti", "Artista": "Jhay Cortez, Bad Bunny", "Sello 1": "The Orchard",
                   "% Sello 1": 0.5, "Sello 2": "Universal", "% Sello 2": 0.5}]).to_excel(lista, index=False)
    df = _fuente([(1, "A", "d", 4e6, "Dakiti", "Jhay Cortez & Bad Bunny"),   # "&" en vez de ","
                  (2, "B", "d", 6e6, "Otra", "x")]
                 + [(p, f"R{p}", "d", 0.0, "Relleno", "r") for p in range(3, 11)])
    df["Sello"] = ["The Orchard", "Sony"] + ["Independientes"] * 8
    df = bmat_calculo.marcar_titularidad(df, bmat_calculo.cargar_titularidad(lista))
    assert df.loc[0, "Titularidad compartida"] == "The Orchard 50% / Universal 50%"
    r = bmat_calculo.calcular_bandas(df, "CAM").set_index("label_group")
    assert r.loc["The Orchard", "streams_millones"] == pytest.approx(2.0)
    assert r.loc["Universal", "streams_millones"] == pytest.approx(2.0)
    assert r.loc["The Orchard", "tracks"] == 1 and r.loc["Universal", "tracks"] == 0


def test_titularidad_rechaza_sellos_desconocidos(tmp_path):
    lista = tmp_path / "tit.xlsx"
    pd.DataFrame([{"Track": "x", "Artista": "y", "Sello 1": "Orchard", "% Sello 1": 0.5,
                   "Sello 2": "Universal M", "% Sello 2": 0.5}]).to_excel(lista, index=False)
    with pytest.raises(ValueError, match="Orchard"):
        bmat_calculo.cargar_titularidad(lista)


def test_streams_catalogo_por_sello_sin_el_error_de_la_plantilla():
    df = _fuente([(1, "A", "d", 10.0, "t", "a"), (2, "B", "d", 20.0, "t", "a"),
                   (3, "C", "d", 30.0, "t", "a"), (4, "D", "d", 40.0, "t", "a")])
    df[config.BMAT_COL_LANZAMIENTO] = ["2020-05-01", "2025-01-01", "2023-12-31", None]
    df["Sello"] = ["Warner", "Warner", "Virgin", "Warner"]
    t = bmat_calculo.streams_catalogo(df)
    assert list(t.columns) == ["Catalogo", "FrontLine"]
    # 10 de catálogo + 40 del track sin fecha, que cuenta como catálogo
    assert t.loc["Warner", "Catalogo"] == 50.0 and t.loc["Warner", "FrontLine"] == 20.0
    assert t.loc["Virgin", "Catalogo"] == 30.0          # el corte es inclusive
    assert t.loc["Sony"].sum() == 0.0


def test_semaforo_del_resumen():
    verde = bmat_calculo._semaforo({"Universal": 14, "Sony": 9, "Warner": 13})
    assert list(verde) == ["Universal"] and verde["Universal"] == bmat_calculo._VERDE
    empate = bmat_calculo._semaforo({"Universal": 4, "Ingrooves": 4, "Sony": 1})
    assert empate == {"Universal": bmat_calculo._AMARILLO, "Ingrooves": bmat_calculo._AMARILLO}
    rojo = bmat_calculo._semaforo({"Universal": 15, "Sony": 24, "Warner": 22, "Virgin": 3})
    assert set(rojo) == {"Sony", "Warner"} and rojo["Sony"] == bmat_calculo._ROJO
    assert bmat_calculo._semaforo({"Universal": 0, "Sony": 0}) == {}   # banda vacía, sin color


def test_resumen_pinta_el_semaforo(tmp_path):
    df = _fuente([(p, f"I{p}", "d", 1000.0 * (11 - p), f"T{p}", "a") for p in range(1, 11)])
    df["Sello"] = ["Universal"] * 4 + ["Sony"] * 6        # Sony le gana a Universal
    df["Disqueras"] = df["Sello"]
    bandas = bmat_calculo.calcular_bandas(df, "CO")
    ruta = bmat_calculo.escribir_intermedio(tmp_path / "x.xlsx", "CO", 2026, 36, df, bandas)
    ws = openpyxl.load_workbook(ruta)["Resumen"]
    assert ws.cell(row=5, column=2).value == 4                     # TRACKS Universal Top 10
    assert ws.cell(row=5, column=2).fill.fill_type is None          # Universal no lidera: sin color
    assert ws.cell(row=8, column=2).value == 6                      # Sony
    assert ws.cell(row=8, column=2).fill.fgColor.rgb.endswith("FFC7CE")
    assert ws.cell(row=13, column=2).value == 10                    # Total, sin color
    assert ws.cell(row=13, column=2).fill.fill_type is None


# --- clasificación ------------------------------------------------------------------

def test_clasificar_conocidos_regla_y_distribuidora_desconocida():
    tabla = _tabla_clasif([("K1", "Universal Music Group", "Dist A"),
                           ("K2", "Virgin", "Dist A"), ("K3", "Virgin", "Dist A"),
                           ("K4", "Ingrooves", "Dist B")])
    df = _fuente([(1, "K1", "Dist A", 1, "t", "a"),
                  (2, "NUEVO1, K4", "Otra", 1, "t", "a"),     # basta un ISRC conocido
                  (3, "NUEVO2", "Dist A", 1, "t", "a"),       # regla: lo más común de Dist A
                  (4, "NUEVO3", "oneRPM", 1, "t", "a")])      # distribuidora nunca vista
    c = clasif.clasificar(df, tabla)
    assert list(c["Disqueras"]) == ["Universal Music Group", "INgrooves", "Virgin", "oneRPM"]
    assert list(c["Sello"]) == ["Universal", "Ingrooves", "Virgin", "Independientes"]
    assert list(c["Clasificacion"]) == ["conocido", "conocido", "nuevo", "nuevo"]
    assert c.loc[2, "_confianza"] == pytest.approx(2 / 3)


def test_nuevos_se_guardan_y_la_misma_semana_los_sigue_listando():
    clasif.sembrar()  # crea la tabla vacía
    df = _fuente([(1, "N1", "oneRPM", 5, "t", "a")])
    c = clasif.clasificar(df, clasif.cargar(), "semana 36 de 2026")
    clasif.guardar_nuevos(c, "semana 36 de 2026")
    # Otra corrida de la semana 36: sigue siendo nuevo (para la lista)...
    assert clasif.clasificar(df, clasif.cargar(), "semana 36 de 2026")["Clasificacion"][0] == "nuevo"
    # ...y la semana 37 ya lo conoce.
    assert clasif.clasificar(df, clasif.cargar(), "semana 37 de 2026")["Clasificacion"][0] == "conocido"


def test_aplicar_revision_corrige_y_manda_sobre_la_regla(tmp_path):
    clasif.sembrar()
    df = _fuente([(1, "N1, N2", "Dist X", 5, "Canción", "Artista")])
    c = clasif.clasificar(df, clasif.cargar(), "semana 36 de 2026")
    clasif.guardar_nuevos(c, "semana 36 de 2026")
    lista = clasif.lista_para_revisar({"CO": c})
    ruta = clasif.escribir_lista_revision(lista, tmp_path / "revisar.xlsx", "t")

    wb = openpyxl.load_workbook(ruta)
    ws = wb[clasif.HOJA_REVISION]
    encabezado = [x.value for x in ws[4]]
    ws.cell(row=5, column=encabezado.index(clasif.COLUMNA_CORRECCION) + 1, value="virgin")
    wb.save(ruta)

    assert clasif.aplicar_revision(ruta) == 1
    tabla = clasif.cargar().set_index("isrc")
    assert tabla.loc["N1", "disqueras"] == "Virgin" and tabla.loc["N2", "disqueras"] == "Virgin"
    assert tabla.loc["N1", "origen"] == clasif.ORIGEN_REVISION
    # Una nueva corrida de la semana ya no lo trae como nuevo, y no pisa la revisión.
    c2 = clasif.clasificar(df, clasif.cargar(), "semana 36 de 2026")
    assert c2["Sello"][0] == "Virgin" and c2["Clasificacion"][0] == "conocido"
    clasif.guardar_nuevos(c2, "semana 36 de 2026")
    assert clasif.cargar().set_index("isrc").loc["N1", "disqueras"] == "Virgin"


def test_lista_para_revisar_une_mercados_y_prioriza():
    tabla = _tabla_clasif([(f"K{i}", "Virgin", "Dist A") for i in range(9)] +
                          [("J1", "Virgin", "Dist B"), ("J2", "Sony Music Entertainment", "Dist B")])
    co = clasif.clasificar(_fuente([(500, "N1", "Dist A", 3, "t1", "a"),
                                    (900, "N2", "Dist B", 1, "t2", "a")]), tabla)
    pe = clasif.clasificar(_fuente([(150, "N1", "Dist A", 2, "t1", "a")]), tabla)
    lista = clasif.lista_para_revisar({"CO": co, "PE": pe})
    assert len(lista) == 2
    n1 = lista[lista["ISRCs"] == "N1"].iloc[0]
    assert n1["Mejor posición"] == 150 and n1["Mercados"] == "COL, PE" and n1["Streams (total)"] == 5
    assert n1["Prioridad"] == "Alta"                              # Top 200
    n2 = lista[lista["ISRCs"] == "N2"].iloc[0]
    assert n2["Prioridad"] == "Alta" and n2["Confianza regla"] == 0.5   # regla dudosa


# --- histórico y reportes -------------------------------------------------------------

def test_semilla_es_idempotente_y_no_pisa_lo_calculado(tmp_path):
    filas = _banda_hist("CO", 2026, 35, 10, {"Universal": (6, 6.0), "Sony": (4, 4.0)})
    assert _sembrar(tmp_path, filas) == len(config.BMAT_LABELS)
    assert _sembrar(tmp_path, filas) == 0
    assert bmat_calculo.ultima_semana("CO") == (2026, 35)


def test_guardar_semana_reemplaza():
    b = pd.DataFrame(_banda_hist("CO", 2026, 36, 10, {"Universal": (10, 1.0)}))
    bmat_calculo.guardar_semana(2026, 36, "CO", b)
    bmat_calculo.guardar_semana(2026, 36, "CO", b)
    df = history.cargar_bmat_weekly()
    assert len(df) == len(config.BMAT_LABELS)


def test_bmat_no_se_borra_al_recargar_spotify(tmp_path):
    _sembrar(tmp_path, _banda_hist("CO", 2026, 35, 10, {"Universal": (10, 1.0)}))
    assert "bmat_weekly" not in history.TABLAS_HISTORICO
    history.vaciar_historico()
    assert not history.cargar_bmat_weekly().empty


def test_calendario_semana_53_y_separador_de_anio():
    df = pd.DataFrame({"anio": [2020, 2020, 2021], "semana": [1, 53, 2]})
    semanas = bmat.calendario(df)
    assert semanas[0] == (2020, 1) and (2020, 53) in semanas and semanas[-1] == (2021, 2)
    cols = bmat._columnas(semanas, 5)
    assert cols[(2021, 1)] == cols[(2020, 53)] + 2          # una columna vacía entre años
    df2 = pd.DataFrame({"anio": [2021, 2022], "semana": [4, 1]})
    assert len(bmat.calendario(df2)) == 52 + 1


def test_reporte_familia_a(tmp_path):
    filas = (_banda_hist("CO", 2026, 35, 10, {"Universal": (3, 3.0), "Sony": (7, 7.0)})
             + _banda_hist("CO", 2026, 35, 50, {"Universal": (30, 6.0), "Sony": (20, 4.0)})
             + _banda_hist("CO", 2026, 34, 10, {"Universal": (5, 5.0), "Sony": (5, 5.0)}))
    _sembrar(tmp_path, filas)
    salida = bmat.generar_reporte("COL", tmp_path / "col.xlsx")
    wb = openpyxl.load_workbook(salida)
    hoja_ms, hoja_det = config.BMAT_REPORTES["COL"]["hojas"]["CO"]
    assert wb.sheetnames == [hoja_ms, hoja_det]

    ms = wb[hoja_ms]
    assert ms.cell(row=2, column=5).value == "Sem\n1"
    col35 = 5 + 34
    assert ms.cell(row=2, column=col35).value == "Sem\n35"
    assert [ms.cell(row=4 + k, column=3).value for k in range(8)] == config.BMAT_LABELS
    assert ms.cell(row=4, column=col35).value == pytest.approx(0.3)          # Universal Top 10
    assert ms.cell(row=4 + 3, column=col35).value == pytest.approx(0.7)      # Sony
    assert ms.cell(row=13, column=col35).value == pytest.approx(0.6)         # Top 50
    assert ms.cell(row=13, column=col35 - 1).value == "NA"                   # sem 34 sin Top 50
    assert ms.cell(row=4, column=5).value == "NA"                            # sem 1 sin datos
    assert ms.conditional_formatting                                          # semáforo
    assert ms.cell(row=4 + 9 * 7, column=2).value == "TOP\n10.000"

    det = wb[hoja_det]
    col35d = 4 + 34
    assert det.cell(row=4, column=col35d).value == "Top 10"
    assert det.cell(row=6, column=col35d).value == 3                          # TRACKS Universal
    assert det.cell(row=6 + 8, column=col35d).value == 10                     # Total
    assert det.cell(row=16, column=col35d).value == pytest.approx(3.0)        # Streams
    assert det.cell(row=26, column=col35d).value == pytest.approx(0.3)        # %
    assert det.cell(row=37, column=1).value == config.BMAT_FUENTE
    assert det.cell(row=39, column=col35d).value == 30                        # Top 50 TRACKS


def test_reporte_familia_b_tiene_16_hojas_con_su_titulo(tmp_path):
    filas = []
    for m in config.BMAT_REPORTES["CAM"]["hojas"]:
        filas += _banda_hist(m, 2026, 35, 10, {"Universal": (6, 6.0), "Sony": (4, 4.0)})
    _sembrar(tmp_path, filas)
    wb = openpyxl.load_workbook(bmat.generar_reporte("CAM", tmp_path / "cam.xlsx"))
    assert len(wb.sheetnames) == 16
    assert wb.sheetnames[:2] == ["MS CAM", "MS CR"] and wb.sheetnames[8] == "CAM"
    assert "PN" in wb.sheetnames and "MS PN" in wb.sheetnames
    assert wb["HN"]["A4"].value == "TOP 3.000 APDIF HONDURAS"                 # no "Rep Dominicana"
    ms = wb["MS CR"]
    assert ms["B3"].value == "% STREAMS - \nCosta Rica"
    assert ms.cell(row=3, column=5 + 34).value == 35
    assert ms.cell(row=5, column=5 + 34).value == pytest.approx(0.6)
    det = wb["CR"]
    assert det.cell(row=4, column=4).value == "Week"
    assert det.cell(row=8, column=4 + 34).value == 6                           # TRACKS Universal


# --- proceso completo ------------------------------------------------------------------

def _escribir_wk(carpeta, nombre, n, isrc_prefijo, distribuidora="Universal Music Group"):
    filas = [(p, f"{isrc_prefijo}{p}", distribuidora if p % 2 else "oneRPM", 1000.0 * (n + 1 - p),
              f"Track {p}", f"Artista {p}") for p in range(1, n + 1)]
    df = _fuente(filas)
    df.loc[0, config.BMAT_COL_MOVIMIENTO] = "NUEVO"
    df.to_excel(carpeta / nombre, index=False)


def test_generar_semana_completa_y_revision(tmp_path, monkeypatch):
    carpeta = tmp_path / "wk36"
    carpeta.mkdir()
    _escribir_wk(carpeta, "WK36-CO.xlsx", 250, "CO")
    _escribir_wk(carpeta, "WK36-CR.xlsx", 120, "CR")
    # Una clasificación conocida chica, como la semilla real: con ella la
    # regla de cada distribuidora no cambia por una sola corrección.
    semilla = tmp_path / "clasif.csv.gz"
    pd.DataFrame([{"isrc": f"S{i}", "disqueras": d, "track": "", "artista": "", "disquera_original": "",
                   "distribuidora_original": d} for i, d in enumerate(["Universal Music Group"] * 5 + ["oneRPM"] * 5)]
                 ).to_csv(semilla, index=False, compression="gzip")
    clasif.sembrar(semilla)

    monkeypatch.setattr(config, "BMAT_REPORTES_ACTIVOS", ["COL"])
    r = bmat_proceso.generar_semana(carpeta, salida=tmp_path / "out", hoy=date(2026, 9, 28))
    assert (r.anio, r.semana) == (2026, 36)
    assert r.mercados == ["CO"]                                # el WK36-CR se ignora
    assert [p.name for p in r.reportes] == ["MS BMAT COL a Sem 36 de 2026.xlsx"]
    assert [p.name for p in r.intermedios] == ["Top 10000 BMAT COL a sem 36 de 2026.xlsx"]
    assert r.avisos == ["WK36-CO.xlsx trae 250 tracks (se esperaban 10000): las bandas más grandes salen NA."]
    assert r.tracks_nuevos == 250                              # tabla casi vacía: todos nuevos
    assert set(p.name for p in r.carpeta_salida.iterdir()) == {
        "MS BMAT COL a Sem 36 de 2026.xlsx", "Top 10000 BMAT COL a sem 36 de 2026.xlsx",
        "Revisar clasificacion BMAT sem 36 de 2026.xlsx"}

    hist = history.cargar_bmat_weekly()
    co = hist[(hist["country_code"] == "CO") & (hist["banda"] == 10)].set_index("label_group")
    assert co.loc["Universal", "tracks"] == 5 and co.loc["Independientes", "tracks"] == 5

    inter = openpyxl.load_workbook(r.carpeta_salida / "Top 10000 BMAT COL a sem 36 de 2026.xlsx")
    assert inter.sheetnames == ["Streams Catalogo", "Resumen", "TOP 200 Nuevos",
                                "Tracks Independientes", "Archivo Base"]
    base = inter["Archivo Base"]
    assert base.max_row == 251 and base["A1"].value == config.BMAT_COL_POSICION   # 250 tracks
    assert [c.value for c in base[1]][-3:] == ["Sello", "Clasificacion", "Titularidad compartida"]
    assert inter["TOP 200 Nuevos"]["C5"].value == "Track 1"

    # El usuario corrige el track 1 (Universal -> Virgin) y vuelve a correr.
    wb = openpyxl.load_workbook(r.lista_revision)
    ws = wb[clasif.HOJA_REVISION]
    enc = [x.value for x in ws[4]]
    for fila in range(5, ws.max_row + 1):
        if ws.cell(row=fila, column=enc.index("ISRCs") + 1).value == "CO1":
            ws.cell(row=fila, column=enc.index(clasif.COLUMNA_CORRECCION) + 1, value="Virgin")
    wb.save(r.lista_revision)
    r2 = bmat_proceso.generar_semana(carpeta, salida=tmp_path / "out", revision=r.lista_revision,
                                     hoy=date(2026, 9, 28))
    assert r2.correcciones == 1 and r2.tracks_nuevos == 249
    hist = history.cargar_bmat_weekly()
    co = hist[(hist["country_code"] == "CO") & (hist["banda"] == 10)].set_index("label_group")
    assert co.loc["Universal", "tracks"] == 4 and co.loc["Virgin", "tracks"] == 1
    assert len(hist[(hist["country_code"] == "CO")]) == 4 * len(config.BMAT_LABELS)  # 10/50/100/200


def test_intermedio_puede_llevar_las_hojas_opcionales(tmp_path):
    df = _fuente([(p, f"I{p}", "d", 1000.0, f"T{p}", "a") for p in range(1, 11)])
    df["Disqueras"], df["Sello"] = "Universal Music Group", "Universal"
    bandas = bmat_calculo.calcular_bandas(df, "CO")
    ruta = bmat_calculo.escribir_intermedio(tmp_path / "x.xlsx", "CO", 2026, 36, df, bandas,
                                            config.BMAT_HOJAS_INTERMEDIO_DISPONIBLES)
    assert openpyxl.load_workbook(ruta).sheetnames[-1] == "TOP 50 - Posiciones UMG"


def test_por_defecto_se_generan_los_cuatro_reportes(tmp_path):
    carpeta = tmp_path / "wk36"
    carpeta.mkdir()
    for m, n in [("CO", 20), ("PE", 20), ("EC", 20), ("CAM", 20), ("PA", 20)]:
        _escribir_wk(carpeta, f"WK36-{m}.xlsx", n, m)
    r = bmat_proceso.generar_semana(carpeta, salida=tmp_path / "out", hoy=date(2026, 9, 28))
    assert r.mercados == ["CO", "PE", "EC", "CAM", "PA"]
    assert [p.name for p in r.reportes] == [
        "MS BMAT COL a Sem 36 de 2026.xlsx", "MS BMAT Peru a Sem 36 de 2026.xlsx",
        "MS BMAT Ecuador a Sem 36 de 2026.xlsx", "MS BMAT CAM a Sem 36 de 2026.xlsx"]
    assert "Top 3000 BMAT PN a sem 36 de 2026.xlsx" in [p.name for p in r.intermedios]
    assert any("faltó el WK36" in a and "Costa Rica" in a for a in r.avisos)
