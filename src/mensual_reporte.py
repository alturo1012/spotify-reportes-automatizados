"""Escribe el Excel del reporte mensual "Market Share Spotify Latam".

Tres tipos de hoja, igual que el archivo real:

- **"XX-Det"** (una por país): el cálculo. Por cada banda, tres bloques de 7
  filas -- `Tracks TOP N`, `Streams TOP N` y `Streams (%) TOP N`.
- **hojas de país** ("CO", "PE"...): solo el % de streams, con los sellos en
  el orden de presentación y las 5 bandas.
- **"Resumen"**: el % de Universal de cada país, en las bandas 10/50/100/200.

Las columnas (y en el Resumen, las filas) son el calendario: los meses de
cada año más sus cierres Q1..Q4, H1 y el año.

REGLAS DE LOS CIERRES (Q / H1 / año)
------------------------------------
Salen de la plantilla real y se comprobaron celda por celda contra ella
(2019 en adelante, 100% de coincidencia; 2018, que fue el primer año y tiene
valores escritos a mano, no cuadra siempre):

- **streams**: suma de los meses del período.
- **tracks**: promedio simple de los meses del período.
- **%**: se recalcula sobre los streams del período (streams del sello /
  streams de los 7 sellos), no es el promedio de los % mensuales.
- **el año, en tracks**, es la excepción: en la plantilla NO es el promedio
  de los 12 meses sino el promedio de tres columnas, `(H1 + Q3 + Q4) / 3`.
  Da distinto que el promedio de los meses (en Colombia 2018: 3,74 contra
  3,41). Se replica tal cual para que el reporte generado dé lo mismo que el
  que venían haciendo a mano. DECIDIDO el 30/09/2026: se deja así, no se
  "corrige" al promedio de los doce meses.
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import config


AZUL = "FF002060"
GRIS = "FFD9D9D9"
BLANCO = "FFFFFFFF"

FMT_TRACKS = '#,##0_ ;[Red]\\-#,##0\\ '
FMT_STREAMS = '#,##0.00_ ;[Red]\\-#,##0.00\\ '
FMT_PCT = "0%"

_BORDE = Border(*[Side(style="thin", color="FFBFBFBF")] * 4)
# Línea negra fina: la que usa la plantilla para encajonar los encabezados y
# para separar un país del siguiente.
_LINEA = Side(style="thin", color="FF000000")

# Anchos de columna del "Resumen", tal cual los trae la plantilla.
ANCHO_IZQUIERDA = 1.45     # columna A, un margen angosto
ANCHO_MESES = 11.0         # columna B, las etiquetas de mes
ANCHO_SEPARADOR = 1.0      # la columna entre un país y el siguiente
ANCHO_BANDA = 4.54         # cada una de las cuatro columnas de banda


def _pintar(celda, colores) -> None:
    """Pinta una celda con un par (relleno, color de letra) de config."""
    if colores is None:
        return
    fondo, texto = colores
    celda.fill = PatternFill("solid", start_color=fondo, end_color=fondo)
    celda.font = Font(size=celda.font.sz or 10, color=texto)


def color_universal(valor):
    """El color de una celda de % de Universal en el "Resumen", con los
    mismos umbrales que trae el formato condicional de la plantilla.

    El objetivo de participación de Universal es el 30%. Los cortes están
    puestos de forma que el color cuadre con el número redondeado que se ve:
    31% o más verde, 30% amarillo, 29% o menos rojo. Por encima del 40% se
    pinta de celeste, que en la plantilla es una regla aparte y manda sobre
    el verde.
    """
    if valor is None:
        return None
    if valor > config.MENSUAL_UMBRAL_DESTACADO:
        return config.MENSUAL_COLOR_DESTACADO
    if valor > config.MENSUAL_UMBRAL_VERDE:
        return config.MENSUAL_COLOR_VERDE
    if valor < config.MENSUAL_UMBRAL_ROJO:
        return config.MENSUAL_COLOR_ROJO
    return config.MENSUAL_COLOR_AMARILLO


def colores_del_bloque(valores_por_sello: dict) -> dict:
    """El color de cada sello dentro de un bloque de una hoja de país.

    Es la regla de la plantilla, que ahí vive como formato condicional:
    Universal va en VERDE cuando es el sello más alto del bloque, y cualquier
    sello que le gane a Universal va en ROJO. Los demás quedan sin color.
    """
    universal = valores_por_sello.get("Universal")
    if universal is None:
        return {}
    otros = {s: v for s, v in valores_por_sello.items() if s != "Universal" and v is not None}
    por_encima = {s: config.MENSUAL_COLOR_ROJO for s, v in otros.items() if v > universal}
    if por_encima:
        return por_encima
    return {"Universal": config.MENSUAL_COLOR_VERDE}

# Etiquetas de los tres bloques de cada banda en las hojas Det.
BLOQUES = [
    ("tracks", "Tracks\nTOP\n{banda}", FMT_TRACKS),
    ("streams", "Streams\nTOP\n{banda}", FMT_STREAMS),
    ("pct", "Streams\n(%)\nTOP {banda}", FMT_PCT),
]

# Cierres: qué meses entran en cada uno y en qué orden van las columnas.
CIERRES = {"Q1": (1, 2, 3), "Q2": (4, 5, 6), "Q3": (7, 8, 9), "Q4": (10, 11, 12),
           "H1": (1, 2, 3, 4, 5, 6)}
SECUENCIA = [1, 2, 3, "Q1", 4, 5, 6, "Q2", "H1", 7, 8, 9, "Q3", 10, 11, 12, "Q4", "YEAR"]


def calendario(desde: tuple, hasta: tuple) -> list:
    """La secuencia de columnas (o de filas, en el Resumen) entre dos meses.

    Cada entrada es `(anio, clave)`, donde clave es un número de mes o el
    nombre de un cierre. El primer año arranca en el mes de `desde` (la
    plantilla real empieza en mayo de 2017, sin enero-abril ni Q1), y el
    último año se dibuja completo aunque todavía no tenga datos: así las
    columnas no se mueven de lugar cada mes.
    """
    anio_desde, mes_desde = desde
    salida = []
    for anio in range(anio_desde, hasta[0] + 1):
        for clave in SECUENCIA:
            if anio == anio_desde:
                if isinstance(clave, int) and clave < mes_desde:
                    continue
                if clave in CIERRES and max(CIERRES[clave]) < mes_desde:
                    continue
            salida.append((anio, clave))
    return salida


def etiqueta(anio: int, clave) -> str:
    if isinstance(clave, int):
        return config.MESES_ES_ABREV[clave].capitalize()
    if clave == "YEAR":
        return f"Total\n{anio}"
    if clave == "H1":
        return f"H1\n{anio}"
    return clave


class Valores:
    """Los números de un país: mes a mes y por banda, y los cierres
    calculados a partir de ellos."""

    def __init__(self, filas, hasta: tuple = None, cierres=None):
        # (banda, mes) -> {label: (tracks, streams)}
        self.por_mes = {}
        # Meses que recalculó el programa desde sus semanas: si alguno cae
        # dentro de un período, el cierre se recalcula en vez de usar el que
        # venía en la plantilla.
        self.recalculados = set()
        for fila in filas:
            mes = (int(fila["anio"]), int(fila["mes"]))
            self.por_mes.setdefault((int(fila["banda"]), mes), {})[fila["label_group"]] = (
                fila["tracks"], fila["streams_millones"])
            if fila.get("origen") and fila["origen"] != "plantilla":
                self.recalculados.add(mes)
        # (anio, periodo, banda, label) -> (tracks, streams, pct)
        self.cierres = cierres or {}
        self.hasta = hasta
        self._cache = {}

    def cierre_heredado(self, banda: int, anio: int, clave):
        """El cierre tal cual lo traía la plantilla, si ese período no tiene
        ningún mes recalculado por el programa.

        En 2017 y 2018 los cierres de la plantilla no cuadran con sus meses
        (fórmulas viejas: el Q1 de 2018, por ejemplo, se olvidó de enero).
        Recalcularlos daría un archivo "más correcto" pero distinto del que
        el área viene mirando hace años, así que para los períodos que ya
        venían cerrados se respeta el número de la plantilla.
        """
        if not self.cierres or isinstance(clave, int):
            return None
        if any(m in self.recalculados for m in self.meses_de(anio, clave)):
            return None
        return self.cierres.get((anio, clave, banda))

    def cerrado(self, banda: int, anio: int, clave) -> bool:
        """Un cierre (Q, H1, año) solo se escribe cuando el período terminó y
        no le falta ningún mes.

        Lo primero, porque un "Total 2026" con nueve meses sería engañoso (la
        plantilla real hace lo mismo: deja esas columnas en blanco hasta que
        el período termina). Lo segundo, porque si a un trimestre le falta un
        mes -- por ejemplo agosto, cuyas semanas todavía no se cargaron -- el
        Q3 saldría con dos meses y nadie lo notaría al mirarlo.
        """
        if isinstance(clave, int) or self.hasta is None:
            return True
        meses = self.meses_de(anio, clave)
        if max(meses) > self.hasta:
            return False
        dentro = [m for m in meses if m >= config.MENSUAL_INICIO]
        if not dentro:
            return True
        faltan = [m for m in dentro if (banda, m) not in self.por_mes]
        return not faltan or len(faltan) == len(dentro)

    def meses_de(self, anio: int, clave) -> list:
        if isinstance(clave, int):
            return [(anio, clave)]
        if clave == "YEAR":
            return [(anio, m) for m in range(1, 13)]
        return [(anio, m) for m in CIERRES[clave]]

    def _bruto(self, banda: int, anio: int, clave):
        """(tracks promedio, streams sumados) por sello, o None si no hay
        ningún mes con datos."""
        presentes = [self.por_mes[(banda, m)] for m in self.meses_de(anio, clave)
                     if (banda, m) in self.por_mes]
        if not presentes:
            return None
        salida = {}
        for label in config.LABEL_GROUPS_MS:
            tracks = [d[label][0] for d in presentes if label in d and d[label][0] is not None]
            streams = [d[label][1] for d in presentes if label in d and d[label][1] is not None]
            salida[label] = (sum(tracks) / len(tracks) if tracks else None,
                             sum(streams) if streams else None)
        return salida

    def celda(self, banda: int, anio: int, clave, tipo: str, label: str):
        """El valor que va en una celda, o None si ese mes/cierre no tiene
        datos todavía."""
        # El cierre heredado de la plantilla manda: ese período ya lo cerró
        # el área en su momento, con los datos que tenía.
        if not self.cerrado(banda, anio, clave) and not self.cierre_heredado(banda, anio, clave):
            return None
        datos = self._cache.get((banda, anio, clave))
        if datos is None:
            datos = self._calcular(banda, anio, clave)
            self._cache[(banda, anio, clave)] = datos
        if datos is False:
            return None
        return datos[tipo].get(label)

    def _calcular(self, banda: int, anio: int, clave):
        heredado = self.cierre_heredado(banda, anio, clave)
        if heredado:
            return {tipo: {lab: v[i] for lab, v in heredado.items()}
                    for i, tipo in enumerate(("tracks", "streams", "pct"))}
        bruto = self._bruto(banda, anio, clave)
        if bruto is None:
            return False
        tracks = {lab: v[0] for lab, v in bruto.items()}
        streams = {lab: v[1] for lab, v in bruto.items()}
        if clave == "YEAR":
            # El año, en tracks, es (H1 + Q3 + Q4) / 3 -- ver el docstring.
            tracks = {}
            for lab in config.LABEL_GROUPS_MS:
                trozos = [self.celda(banda, anio, k, "tracks", lab) for k in ("H1", "Q3", "Q4")]
                trozos = [t for t in trozos if t is not None]
                tracks[lab] = sum(trozos) / 3 if len(trozos) == 3 else (
                    sum(trozos) / len(trozos) if trozos else None)
        total = sum(v for v in streams.values() if v)
        pct = {lab: (streams[lab] / total if total and streams[lab] is not None else None)
               for lab in config.LABEL_GROUPS_MS}
        return {"tracks": tracks, "streams": streams, "pct": pct}


def _encabezado(ws, fila: int, columnas: list, col_inicio: int, con_fecha: bool) -> None:
    for i, (anio, clave) in enumerate(columnas):
        celda = ws.cell(fila, col_inicio + i)
        celda.value = etiqueta(anio, clave)
        celda.font = Font(bold=True, size=9, color=BLANCO)
        celda.fill = PatternFill("solid", start_color=AZUL, end_color=AZUL)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_inicio + i)].width = 9.6


def hoja_det(wb, pais: str, valores: Valores, columnas: list) -> None:
    ws = wb.create_sheet(config.MENSUAL_HOJA_DET.format(sigla=pais))
    ws.sheet_state = "visible"
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 18
    _encabezado(ws, 1, columnas, 4, con_fecha=True)

    fila = 3
    for banda in config.MENSUAL_BANDAS:
        for tipo, titulo, formato in BLOQUES:
            celda = ws.cell(fila, 1)
            celda.value = titulo.format(banda=banda)
            celda.font = Font(bold=True, size=9)
            celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.merge_cells(start_row=fila, start_column=1, end_row=fila + 6, end_column=1)
            for j, label in enumerate(config.LABEL_GROUPS_MS):
                ws.cell(fila + j, 2).value = config.MENSUAL_NOMBRES_DET[tipo][label]
                ws.cell(fila + j, 2).font = Font(size=9)
                for i, (anio, clave) in enumerate(columnas):
                    c = ws.cell(fila + j, 4 + i)
                    c.value = valores.celda(banda, anio, clave, tipo, label)
                    c.number_format = formato
                    c.font = Font(size=9)
            fila += 7
        fila += 1
    ws.freeze_panes = "D2"


def hoja_pais(wb, pais: str, nombre: str, valores: Valores, columnas: list) -> None:
    ws = wb.create_sheet(config.MENSUAL_HOJAS_PAIS[pais])
    ws.column_dimensions["A"].width = 9.5
    ws.column_dimensions["B"].width = 20.6
    titulo = ws.cell(2, 1)
    titulo.value = f"% Streams - {nombre}"
    titulo.font = Font(bold=True, size=11, color=BLANCO)
    titulo.fill = PatternFill("solid", start_color=AZUL, end_color=AZUL)
    _encabezado(ws, 2, columnas, 3, con_fecha=False)

    fila = 4
    for banda in config.MENSUAL_BANDAS:
        celda = ws.cell(fila, 1)
        celda.value = f"TOP\n{banda}"
        celda.font = Font(bold=True, size=10)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila + 6, end_column=1)
        # Los colores se deciden por bloque (banda x mes), comparando los
        # siete sellos entre sí -- ver `colores_del_bloque`.
        colores = {
            (anio, clave): colores_del_bloque(
                {lab: valores.celda(banda, anio, clave, "pct", lab)
                 for lab in config.MENSUAL_ORDEN_PAIS})
            for (anio, clave) in columnas
        }
        for j, label in enumerate(config.MENSUAL_ORDEN_PAIS):
            ws.cell(fila + j, 2).value = config.MENSUAL_NOMBRES_PAIS[label]
            ws.cell(fila + j, 2).font = Font(size=10)
            for i, (anio, clave) in enumerate(columnas):
                c = ws.cell(fila + j, 3 + i)
                c.value = valores.celda(banda, anio, clave, "pct", label)
                c.number_format = FMT_PCT
                c.font = Font(size=10)
                c.border = _BORDE
                _pintar(c, colores[(anio, clave)].get(label))
        fila += 8
    ws.freeze_panes = "C3"


def hoja_resumen(wb, valores_por_pais: dict, columnas: list) -> None:
    ws = wb.create_sheet("Resumen", 0)
    # Anchos exactos de la plantilla. La columna que separa un país del
    # siguiente es angosta a propósito (1 carácter): lo que divide los países
    # son las líneas de los costados, no un hueco. Si se deja el ancho por
    # defecto quedan unos espacios enormes entre bloques.
    ws.column_dimensions["A"].width = ANCHO_IZQUIERDA
    ws.column_dimensions["B"].width = ANCHO_MESES
    ws.column_dimensions["C"].width = ANCHO_SEPARADOR
    titulo = ws.cell(1, 2)
    titulo.value = "Resumen / Market Share UM"
    titulo.font = Font(bold=True, size=12, color=BLANCO)
    titulo.alignment = Alignment(horizontal="center", vertical="center")
    for c in range(2, 8):
        ws.cell(1, c).fill = PatternFill("solid", start_color=AZUL, end_color=AZUL)
    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=7)

    # Encabezados como en la plantilla: sin relleno, negrita negra y
    # centrados, con un recuadro de línea fina. Las líneas verticales de los
    # costados siguen bajando por todas las filas de datos, que es lo que
    # separa visualmente un país del siguiente.
    mes_top = ws.cell(3, 2)
    mes_top.value = "MES /TOP"
    mes_top.font = Font(bold=True, size=9)
    mes_top.alignment = Alignment(horizontal="center", vertical="center")
    mes_top.border = Border(left=_LINEA, right=_LINEA, top=_LINEA, bottom=_LINEA)

    ancho = len(config.MENSUAL_BANDAS_RESUMEN)
    col = 4
    for pais, nombre in config.MENSUAL_PAISES:
        for k in range(ancho):
            c = ws.cell(3, col + k)
            c.border = Border(top=_LINEA, bottom=_LINEA,
                              left=_LINEA if k == 0 else None,
                              right=_LINEA if k == ancho - 1 else None)
        cab = ws.cell(3, col)
        cab.value = nombre
        cab.font = Font(bold=True, size=9)
        cab.alignment = Alignment(horizontal="center", vertical="center")
        ws.merge_cells(start_row=3, start_column=col, end_row=3, end_column=col + ancho - 1)
        for k, banda in enumerate(config.MENSUAL_BANDAS_RESUMEN):
            c = ws.cell(4, col + k)
            c.value = banda
            c.font = Font(bold=True, size=9)
            c.alignment = Alignment(horizontal="center")
            c.border = Border(bottom=_LINEA,
                              left=_LINEA if k == 0 else None,
                              right=_LINEA if k == ancho - 1 else None)
            ws.column_dimensions[get_column_letter(col + k)].width = ANCHO_BANDA
        ws.column_dimensions[get_column_letter(col + ancho)].width = ANCHO_SEPARADOR
        col += ancho + 1

    for i, (anio, clave) in enumerate(columnas):
        fila = 5 + i
        etq = ws.cell(fila, 2)
        # Acá el mes lleva el año pegado ("may-17"), como en la plantilla: la
        # columna es larguísima y sin el año uno se pierde al bajar.
        etq.value = (f"{config.MESES_ES_ABREV[clave]}-{anio % 100:02d}"
                     if isinstance(clave, int) else etiqueta(anio, clave).replace("\n", " "))
        etq.font = Font(size=9, bold=not isinstance(clave, int), color="000000")
        etq.fill = PatternFill("solid", start_color=config.MENSUAL_COLOR_MESES,
                               end_color=config.MENSUAL_COLOR_MESES)
        etq.alignment = Alignment(horizontal="center")
        etq.border = _BORDE
        col = 4
        ultima = len(config.MENSUAL_BANDAS_RESUMEN) - 1
        for pais, _ in config.MENSUAL_PAISES:
            valores = valores_por_pais.get(pais)
            for k, banda in enumerate(config.MENSUAL_BANDAS_RESUMEN):
                c = ws.cell(fila, col + k)
                c.value = valores.celda(banda, anio, clave, "pct", "Universal") if valores else None
                c.number_format = FMT_PCT
                c.font = Font(size=9)
                # Cuadrícula tenue adentro y línea negra en los bordes del
                # bloque, para que el país quede encajonado de arriba abajo.
                c.border = Border(
                    top=_BORDE.top, bottom=_BORDE.bottom,
                    left=_LINEA if k == 0 else _BORDE.left,
                    right=_LINEA if k == ultima else _BORDE.right)
                _pintar(c, color_universal(c.value))
            col += len(config.MENSUAL_BANDAS_RESUMEN) + 1
    _leyenda(ws, 1, 9)
    ws.freeze_panes = "D5"


def _leyenda(ws, fila: int, col_inicio: int) -> None:
    """La leyenda de colores, arriba del cuadro. La plantilla trae los
    umbrales sueltos en unas celdas lejos de la vista (BG2:BL2); acá se
    escriben con su color al lado del título, que se entiende sin tener que
    ir a buscarlos."""
    textos = [
        (config.MENSUAL_COLOR_DESTACADO, "40% o más"),
        (config.MENSUAL_COLOR_VERDE, "31% o más"),
        (config.MENSUAL_COLOR_AMARILLO, "30% (objetivo)"),
        (config.MENSUAL_COLOR_ROJO, "29% o menos"),
    ]
    col = col_inicio
    for colores, texto in textos:
        celda = ws.cell(fila, col)
        celda.value = texto
        celda.alignment = Alignment(horizontal="center")
        celda.border = _BORDE
        _pintar(celda, colores)
        celda.font = Font(size=9, color=colores[1])
        ws.merge_cells(start_row=fila, start_column=col, end_row=fila, end_column=col + 3)
        col += 5


def configurar_impresion(ws, filas_titulo: str = None) -> None:
    """Deja la hoja lista para imprimir: apaisada, márgenes angostos, ajustada
    al ancho de la página y repitiendo los encabezados en cada hoja.

    Sin esto, una hoja con más de cien columnas sale partida en pedazos
    imposibles de leer."""
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.25
    ws.page_margins.top = ws.page_margins.bottom = 0.4
    ws.page_margins.header = ws.page_margins.footer = 0.2
    if filas_titulo:
        ws.print_title_rows = filas_titulo


def _cierres_por_pais(cierres) -> dict:
    """De la tabla de cierres heredados al diccionario que usa `Valores`."""
    salida = {}
    if cierres is None or getattr(cierres, "empty", True):
        return salida
    for fila in cierres.to_dict("records"):
        clave = (int(fila["anio"]), fila["periodo"], int(fila["banda"]))
        salida.setdefault(fila["country_code"], {}).setdefault(clave, {})[fila["label_group"]] = (
            fila["tracks"], fila["streams_millones"], fila["pct_streams"])
    return salida


def generar_reporte(mensual, path, hasta: tuple, desde: tuple = None, cierres=None):
    """Escribe el Excel completo con el histórico `mensual` (el DataFrame de
    `history.cargar_ms_mensual()`), hasta el mes `hasta`.

    `cierres` es lo que devuelve `history.cargar_ms_mensual_cierre()`: los
    Q/H1/año tal cual venían en la plantilla, que se respetan mientras el
    período no tenga ningún mes recalculado (ver `Valores.cierre_heredado`).
    """
    desde = desde or config.MENSUAL_INICIO
    columnas = calendario(desde, hasta)
    datos = mensual[(mensual["anio"] * 100 + mensual["mes"]) <= hasta[0] * 100 + hasta[1]]
    heredados = _cierres_por_pais(cierres)

    valores_por_pais = {}
    for pais, _ in config.MENSUAL_PAISES:
        valores_por_pais[pais] = Valores(
            datos[datos["country_code"] == pais].to_dict("records"), hasta=hasta,
            cierres=heredados.get(pais))

    wb = Workbook()
    wb.remove(wb.active)
    hoja_resumen(wb, valores_por_pais, columnas)
    for pais, nombre in config.MENSUAL_PAISES:
        hoja_pais(wb, pais, nombre, valores_por_pais[pais], columnas)
    for pais, _ in config.MENSUAL_PAISES:
        hoja_det(wb, pais, valores_por_pais[pais], columnas)

    for hoja in wb.worksheets:
        # El Resumen lleva los meses en las filas, así que lo que se repite al
        # imprimir son sus dos filas de encabezado; las otras hojas no.
        configurar_impresion(hoja, "3:4" if hoja.title == "Resumen" else None)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def nombre_reporte(anio: int, mes: int) -> str:
    return config.MENSUAL_ARCHIVO.format(mes=config.MESES_ES_ABREV[mes], anio=anio)
