"""Genera los cuatro reportes finales BMAT a partir del histórico guardado:
"MS BMAT COL", "MS BMAT Peru", "MS BMAT Ecuador" y "MS BMAT CAM".

El histórico (tabla bmat_weekly) trae, por mercado, semana, banda y sello:
tracks, streams (millones) y % de streams. Se sembró con los reportes
reales de la semana 35 de 2026 y cada semana nueva la agrega
bmat_calculo.py desde el archivo WK de BMAT. Este módulo solo dibuja.

Dos familias de formato, copiadas de los archivos reales (config.BMAT_REPORTES):

- A (COL, PE, EC): una hoja de % ("Market Share ...") y una de detalle
  ("Resumen/Detalle ...": TRACKS, Streams (Millones) y % Streams).
- B (CAM): 16 hojas, por cada uno de los 8 mercados una "MS XX" (%) y una
  "XX" (detalle), con el año y la semana en otras filas y un título por país.

En las dos, las semanas van en columnas, de la semana 1 del primer año hasta
la última guardada, con una columna angosta separando cada año. Una banda
sin datos en una semana (la fuente de ese año era más chica, o no hubo
archivo) sale "NA" en la hoja de % y vacía en el detalle, como en los
originales.

El semáforo de la hoja de % es la misma regla de los originales (y del
Reporte_MS_TOP200), como formato condicional: verde en Universal cuando es
el mayor de su banda, rojo en el sello que le gana a Universal.

Lo que NO se reproduce de los archivos reales, a propósito:
- las hojas "Presentacion", "Comportamiento" y "TOP20 2018" (de 2018 y
  llenas de #REF!);
- sus errores de copiado: títulos con otro país ("PROMUSICA COLOMBIA" en
  Perú y Ecuador, "Rep Dominicana" en Honduras), un bloque "TOP 5.000"
  sobrante en Ecuador, y la hoja de % de Costa Rica y Guatemala con The
  Orchard Top 1.000 de 2021 corrido una semana. Todo sale del detalle.
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import config, history

_AZUL = PatternFill("solid", start_color=config.COLOR_BANNER_BMAT, end_color=config.COLOR_BANNER_BMAT)
_FINO = Side(style="thin")
_BORDE = Border(left=_FINO, right=_FINO, top=_FINO, bottom=_FINO)
_BORDE_DERECHO = Border(right=_FINO)
_BORDE_ETIQUETA = Border(left=_FINO, right=_FINO, top=_FINO)
_CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)
_IZQ_CENTRO = Alignment(horizontal="left", vertical="center")
_DER_CENTRO = Alignment(horizontal="right", vertical="center")
_V_CENTRO = Alignment(vertical="center")
_V_CENTRO_WRAP = Alignment(vertical="center", wrap_text=True)

# Semáforo (mismos colores que el formato condicional de los originales).
_VERDE = (PatternFill("solid", start_color="C6EFCE", end_color="C6EFCE", bgColor="C6EFCE"),
          Font(color="006100"))
_ROJO = (PatternFill("solid", start_color="FFC7CE", end_color="FFC7CE", bgColor="FFC7CE"),
         Font(color="9C0006"))

_FMT_TRACKS = "#,##0_ ;[Red]\\-#,##0\\ "
_FMT_STREAMS = "#,##0.00_ ;[Red]\\-#,##0.00\\ "
_FMT_PCT_DETALLE = "0%"
_FMT_PCT_MS = "0.0%"

_SUBTABLAS = [  # (etiqueta, campo, desplazamiento desde la fila de la banda, formato)
    ("TRACKS", "tracks", 2, _FMT_TRACKS),
    ("Streams \n(Millones)", "streams_millones", 12, _FMT_STREAMS),
    ("% \nStreams", "pct_streams", 22, _FMT_PCT_DETALLE),
]

# Filas y anchos de cada familia, medidos en los archivos reales.
_LAYOUT = {
    "A": {
        "ms": {"fila_anio": 1, "fila_semana": 2, "fila_banda0": 4, "col_banda": 2, "col_sello": 3,
               "col_semana0": 5, "anchos": {1: 3.45, 2: 10.54, 3: 22.45, 4: 1.18}, "ancho_semana": 9.54,
               "alto_semana": 31.0, "altos": {3: 3.75}, "alto_banda": 21.0, "alto_separador": 6.0,
               "fuente_etiquetas": 12, "prefijo_sem": True, "congelar": "D3"},
        "det": {"fila_anio": 1, "fila_semana": 2, "fila_banda0": 4, "col_semana0": 4,
                "anchos": {1: 8.45, 2: 13.54, 3: 2.0}, "ancho_semana": 7.0,
                "altos": {2: 30.0, 5: 3.0}, "congelar": "D5", "rotulo_week": None},
    },
    "B": {
        "ms": {"fila_anio": 2, "fila_semana": 3, "fila_banda0": 5, "col_banda": 2, "col_sello": 3,
               "col_semana0": 5, "anchos": {1: 3.45, 2: 8.54, 3: 19.45, 4: 1.36}, "ancho_semana": 8.0,
               "alto_semana": 31.5, "altos": {2: 18.5, 4: 6.0}, "alto_banda": 24.0, "alto_separador": 10.5,
               "fuente_etiquetas": 12, "prefijo_sem": False, "congelar": "D4"},
        "det": {"fila_anio": 3, "fila_semana": 6, "fila_banda0": 6, "col_semana0": 4,
                "anchos": {1: 9.36, 2: 23.36, 3: 1.63}, "ancho_semana": 6.5,
                "altos": {3: 18.5, 4: 15.5}, "congelar": "D7", "rotulo_week": 4},
    },
}
_ANCHO_SEPARADOR_ANIO = 1.45
_SALTOS_BANDA_DETALLE = [33] + [32] * 6   # entre banda y banda, como el original
_SALTO_BANDA_MS = 9                       # 8 sellos + 1 fila en blanco


# --- Datos ------------------------------------------------------------------

def cargar_historico(mercados=None) -> pd.DataFrame:
    df = history.cargar_bmat_weekly()
    if mercados is not None and not df.empty:
        df = df[df["country_code"].isin(list(mercados))]
    return df


def calendario(df: pd.DataFrame, hasta=None) -> list:
    """Las semanas que van en columnas: [(anio, semana), ...], de la semana
    1 del primer año con datos hasta la última (o `hasta`). Un año cerrado
    tiene 52 semanas, o 53 si los datos llegan a la 53 (2020)."""
    if df.empty:
        return []
    claves = df[["anio", "semana"]].drop_duplicates()
    if hasta is not None:
        claves = claves[claves["anio"] * 100 + claves["semana"] <= hasta[0] * 100 + hasta[1]]
    if claves.empty:
        return []
    ultimo = claves.sort_values(["anio", "semana"]).iloc[-1]
    ultimo_anio, ultima_semana = int(ultimo["anio"]), int(ultimo["semana"])
    maximos = claves.groupby("anio")["semana"].max()
    semanas = []
    for anio in range(int(claves["anio"].min()), ultimo_anio + 1):
        tope = ultima_semana if anio == ultimo_anio else max(52, int(maximos.get(anio, 52)))
        semanas.extend((anio, s) for s in range(1, tope + 1))
    return semanas


def _columnas(semanas: list, col0: int) -> dict:
    """{(anio, semana): columna}, dejando una columna vacía entre años."""
    cols, col, anio_prev = {}, col0, None
    for anio, semana in semanas:
        if anio_prev is not None and anio != anio_prev:
            col += 1
        cols[(anio, semana)] = col
        col += 1
        anio_prev = anio
    return cols


def _separadores(cols: dict) -> list:
    usadas = set(cols.values())
    return [c for c in range(min(usadas), max(usadas) + 1) if c not in usadas] if usadas else []


def _nombre_banda(banda: int) -> str:
    """10 -> "10"; 1000 -> "1.000" (separador de miles, como los originales)."""
    return f"{banda:,}".replace(",", ".")


def nombre_reporte(reporte: str, anio: int, semana: int) -> str:
    return f"{config.BMAT_REPORTES[reporte]['archivo']} a Sem {semana} de {anio}.xlsx"


def _bandas_mercado(df_m: pd.DataFrame, mercado: str) -> list:
    tope = config.BMAT_MERCADOS[mercado]["top"]
    return [b for b in config.BMAT_BANDAS if b <= tope]


def _titulo_mercado(reporte: str, mercado: str) -> str:
    m = config.BMAT_MERCADOS[mercado]
    titulo = config.BMAT_REPORTES[reporte]["titulo"]
    if config.BMAT_REPORTES[reporte]["familia"] == "B":
        return f"TOP {_nombre_banda(m['top'])} {titulo} {m['nombre'].upper()}"
    return f"TOP {_nombre_banda(m['top'])} \n{titulo}"


# --- Encabezados comunes ----------------------------------------------------

def _encabezado_semanas(ws, cols, lay, tamano, prefijo_sem, formato_semana="General"):
    anio_prev = None
    for (anio, semana), col in cols.items():
        if anio != anio_prev:
            c = ws.cell(row=lay["fila_anio"], column=col, value=anio)
            c.font = Font(bold=True, size=tamano)
            c.alignment = Alignment(horizontal="left", vertical="center")
            anio_prev = anio
        c = ws.cell(row=lay["fila_semana"], column=col, value=f"Sem\n{semana}" if prefijo_sem else semana)
        c.font = Font(bold=True, size=tamano, color="FFFFFF")
        c.fill, c.border, c.alignment = _AZUL, _BORDE, _CENTRO
        c.number_format = formato_semana
        ws.column_dimensions[get_column_letter(col)].width = lay["ancho_semana"]
    for col in _separadores(cols):
        ws.column_dimensions[get_column_letter(col)].width = _ANCHO_SEPARADOR_ANIO
    for col, ancho in lay["anchos"].items():
        ws.column_dimensions[get_column_letter(col)].width = ancho
    for fila, alto in lay.get("altos", {}).items():
        ws.row_dimensions[fila].height = alto


def _congelar_y_mostrar_el_final(ws, celda: str, cols: dict, semanas_visibles: int = 12):
    """Congela títulos y deja la vista en las últimas semanas (lo que se
    mira cada semana), no en 2020."""
    ws.freeze_panes = celda
    if cols and ws.sheet_view.pane is not None:
        ultima = max(cols.values())
        primera = min(cols.values())
        inicio = max(primera, ultima - semanas_visibles)
        fila = ws[celda].row
        ws.sheet_view.pane.topLeftCell = f"{get_column_letter(inicio)}{fila}"


# --- Hoja de % (Market Share) ------------------------------------------------

def _escribir_ms(ws, df_m, mercado, familia, titulo, semanas):
    lay = _LAYOUT[familia]["ms"]
    cols = _columnas(semanas, lay["col_semana0"])
    _encabezado_semanas(ws, cols, lay, 12, lay["prefijo_sem"],
                        "General" if lay["prefijo_sem"] else "#,##0_ ;[Red]\\-#,##0\\ ")
    ws.row_dimensions[lay["fila_semana"]].height = lay["alto_semana"]

    c = ws.cell(row=lay["fila_semana"], column=lay["col_banda"], value=titulo)
    c.font, c.alignment = Font(bold=True, size=12), _CENTRO
    ws.merge_cells(start_row=lay["fila_semana"], start_column=lay["col_banda"],
                   end_row=lay["fila_semana"], end_column=lay["col_sello"])
    ws.cell(row=lay["fila_semana"], column=lay["col_sello"]).border = _BORDE_DERECHO

    pct = {(int(r.anio), int(r.semana), int(r.banda), r.label_group): r.pct_streams
           for r in df_m.itertuples(index=False)}
    con_datos = {(a, s, b) for a, s, b, _ in pct}
    n = len(config.BMAT_LABELS)
    primera_col, ultima_col = (min(cols.values()), max(cols.values())) if cols else (lay["col_semana0"],) * 2
    fuente_valor = Font(size=11)
    fuente_sello = Font(size=lay["fuente_etiquetas"])

    for i, banda in enumerate(_bandas_mercado(df_m, mercado)):
        fila0 = lay["fila_banda0"] + i * _SALTO_BANDA_MS
        c = ws.cell(row=fila0, column=lay["col_banda"], value=f"TOP\n{_nombre_banda(banda)}")
        c.font, c.alignment = Font(bold=True, size=lay["fuente_etiquetas"]), _CENTRO
        for k in range(n):
            ws.cell(row=fila0 + k, column=lay["col_banda"]).border = _BORDE
        ws.merge_cells(start_row=fila0, start_column=lay["col_banda"],
                       end_row=fila0 + n - 1, end_column=lay["col_banda"])
        for k, sello in enumerate(config.BMAT_LABELS):
            c = ws.cell(row=fila0 + k, column=lay["col_sello"], value=sello)
            c.font, c.border, c.alignment = fuente_sello, _BORDE, _IZQ_CENTRO
            ws.row_dimensions[fila0 + k].height = lay["alto_banda"]
        ws.row_dimensions[fila0 + n].height = lay["alto_separador"]

        for (anio, semana), col in cols.items():
            hay = (anio, semana, banda) in con_datos
            for k, sello in enumerate(config.BMAT_LABELS):
                if hay:
                    v = pct.get((anio, semana, banda, sello))
                    v = 0.0 if v is None or pd.isna(v) else float(v)
                else:
                    v = "NA"
                c = ws.cell(row=fila0 + k, column=col, value=v)
                c.font, c.border, c.number_format = fuente_valor, _BORDE, _FMT_PCT_MS
                c.alignment = _DER_CENTRO if hay else _CENTRO

        # Semáforo: dos reglas por banda sobre todas las semanas.
        uni = f"{get_column_letter(primera_col)}{fila0}"
        rango_uni = f"{uni}:{get_column_letter(ultima_col)}{fila0}"
        ws.conditional_formatting.add(rango_uni, FormulaRule(
            formula=[f"AND(ISNUMBER({uni}),{uni}>0,{uni}=MAX({get_column_letter(primera_col)}{fila0}:"
                     f"{get_column_letter(primera_col)}{fila0 + n - 1}))"],
            fill=_VERDE[0], font=_VERDE[1]))
        otros = f"{get_column_letter(primera_col)}{fila0 + 1}"
        rango_otros = f"{otros}:{get_column_letter(ultima_col)}{fila0 + n - 1}"
        ws.conditional_formatting.add(rango_otros, FormulaRule(
            formula=[f"AND(ISNUMBER({otros}),ISNUMBER({get_column_letter(primera_col)}${fila0}),"
                     f"{otros}>{get_column_letter(primera_col)}${fila0})"],
            fill=_ROJO[0], font=_ROJO[1]))

    _congelar_y_mostrar_el_final(ws, lay["congelar"], cols)


# --- Hoja de detalle ------------------------------------------------------------

def _filas_banda_detalle(lay, cantidad):
    filas, fila = [], lay["fila_banda0"]
    for i in range(cantidad):
        filas.append(fila)
        if i < len(_SALTOS_BANDA_DETALLE):
            fila += _SALTOS_BANDA_DETALLE[i]
    return filas


def _escribir_detalle(ws, df_m, mercado, familia, titulo, semanas):
    lay = _LAYOUT[familia]["det"]
    cols = _columnas(semanas, lay["col_semana0"])
    _encabezado_semanas(ws, cols, lay, 9, False)
    fuente = Font(size=9)
    negrita = Font(size=9, bold=True)
    blanca = Font(size=9, bold=True, color="FFFFFF")

    if familia == "A":
        c = ws.cell(row=2, column=1, value=titulo)
        c.font, c.alignment = Font(bold=True, size=12), _CENTRO
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=2)
        ws.cell(row=2, column=2).border = _BORDE_DERECHO
    else:
        c = ws.cell(row=4, column=1, value=titulo)
        c.font, c.alignment = Font(bold=True, size=12), _V_CENTRO
        c = ws.cell(row=6, column=1, value="Fuentes: Google Play, Napster, Spotify")
        c.font, c.alignment = fuente, Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.merge_cells(start_row=6, start_column=1, end_row=6, end_column=2)
        for col in cols.values():
            c = ws.cell(row=lay["rotulo_week"], column=col, value="Week")
            c.font, c.fill, c.border, c.alignment = Font(size=8, bold=True, color="FFFFFF"), _AZUL, _BORDE, _CENTRO

    datos = {(int(r.anio), int(r.semana), int(r.banda), r.label_group): r
             for r in df_m.itertuples(index=False)}
    con_datos = {(a, s, b) for a, s, b, _ in datos}
    bandas = _bandas_mercado(df_m, mercado)
    filas_banda = _filas_banda_detalle(lay, len(bandas))

    for i, (banda, fila_banda) in enumerate(zip(bandas, filas_banda)):
        # La fila de la banda: en la familia B, la primera coincide con la de
        # los números de semana (así es el original) y no lleva rótulo.
        if not (familia == "B" and fila_banda == lay["fila_semana"]):
            for col in cols.values():
                c = ws.cell(row=fila_banda, column=col, value=f"Top {_nombre_banda(banda)}")
                c.font, c.fill, c.border, c.alignment = blanca, _AZUL, _BORDE, _CENTRO
        for etiqueta, campo, offset, formato in _SUBTABLAS:
            fila0 = fila_banda + offset
            c = ws.cell(row=fila0, column=1, value=etiqueta)
            c.font, c.border, c.alignment = negrita, _BORDE_ETIQUETA, _V_CENTRO_WRAP
            ws.row_dimensions[fila0].height = 24.0
            for k in range(1, len(config.BMAT_LABELS) + 1):
                ws.cell(row=fila0 + k, column=1).border = Border(left=_FINO, right=_FINO)
            for k, sello in enumerate(config.BMAT_LABELS + ["Total"]):
                c = ws.cell(row=fila0 + k, column=2, value=sello)
                c.font = negrita if sello == "Total" else fuente
                c.border, c.alignment = _BORDE, _V_CENTRO
            ws.cell(row=fila0 + len(config.BMAT_LABELS), column=1).border = Border(
                left=_FINO, right=_FINO, bottom=_FINO)

            for (anio, semana), col in cols.items():
                hay = (anio, semana, banda) in con_datos
                total = 0.0
                for k, sello in enumerate(config.BMAT_LABELS):
                    v = None
                    if hay:
                        r = datos.get((anio, semana, banda, sello))
                        v = getattr(r, campo) if r is not None else None
                        v = None if v is None or pd.isna(v) else float(v)
                        total += v or 0.0
                    c = ws.cell(row=fila0 + k, column=col, value=v)
                    c.font, c.border, c.number_format, c.alignment = fuente, _BORDE, formato, _DER_CENTRO
                c = ws.cell(row=fila0 + len(config.BMAT_LABELS), column=col, value=total if hay else None)
                c.font, c.border, c.number_format, c.alignment = negrita, _BORDE, formato, _DER_CENTRO

        if i == 1:  # la nota de fuente va en la fila de la segunda banda
            c = ws.cell(row=fila_banda, column=1, value=config.BMAT_FUENTE)
            c.font, c.alignment = fuente, Alignment(horizontal="left", vertical="center", wrap_text=True)
            ws.merge_cells(start_row=fila_banda, start_column=1, end_row=fila_banda, end_column=2)
            ws.row_dimensions[fila_banda].height = 24.0

    _congelar_y_mostrar_el_final(ws, lay["congelar"], cols)


# --- Reporte ----------------------------------------------------------------------

def generar_reporte(reporte: str, output_path: Path, hasta=None) -> Path:
    """Escribe uno de los cuatro reportes ("COL", "PE", "EC", "CAM") con el
    histórico guardado, hasta la semana `hasta` = (anio, semana) si se da."""
    conf = config.BMAT_REPORTES[reporte]
    familia = conf["familia"]
    df = cargar_historico(conf["hojas"].keys())
    if hasta is not None and not df.empty:
        df = df[df["anio"] * 100 + df["semana"] <= hasta[0] * 100 + hasta[1]]

    wb = Workbook()
    wb.remove(wb.active)
    hojas_ms, hojas_det = [], []
    for mercado, (hoja_ms, hoja_det) in conf["hojas"].items():
        df_m = df[df["country_code"] == mercado]
        semanas = calendario(df_m, hasta)
        titulo = _titulo_mercado(reporte, mercado)
        ws = wb.create_sheet(hoja_ms)
        titulo_ms = (f"% STREAMS - \n{config.BMAT_MERCADOS[mercado]['nombre']}" if familia == "B" else titulo)
        _escribir_ms(ws, df_m, mercado, familia, titulo_ms, semanas)
        hojas_ms.append(ws)
        hojas_det.append((hoja_det, df_m, mercado, titulo, semanas))
    # Orden de los originales: primero todas las de %, después los detalles.
    for hoja_det, df_m, mercado, titulo, semanas in hojas_det:
        _escribir_detalle(wb.create_sheet(hoja_det), df_m, mercado, familia, titulo, semanas)

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path
