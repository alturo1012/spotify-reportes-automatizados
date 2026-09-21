"""Cálculo semanal BMAT: de la fuente "WK<semana>-<mercado>.xlsx" al
histórico por banda y sello, más el archivo intermedio de cada mercado.

Es la parte mecánica del proceso manual: pegar la fuente en "Archivo Base",
clasificar cada track (ver bmat_clasificacion.py) y dejar que las tablas
dinámicas sumen por banda. Verificado contra los intermedios de la semana
35: con la misma clasificación, CO y PE dan 0 diferencias en 128 valores
cada uno.

Qué se calcula, por mercado y banda (Top 10, 50, 100, 200, 1.000, 3.000,
5.000, 10.000 -- solo las que caben en la fuente):

- tracks: cuántas canciones del Top N son de cada sello;
- streams_millones: sus "Streams con video", en millones;
- pct_streams: streams del sello / streams totales del Top N.

La semana es la del nombre del archivo (WK35 -> 35), no un contador: BMAT
numera sus semanas y los reportes siempre usaron ese número.
"""
import re
import unicodedata
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import bmat_clasificacion as clasif
from . import config, history

SEED_BMAT_CSV = history.SEED_BMAT_CSV

_PATRON_ARCHIVO = re.compile(r"WK\s*(\d{1,2})\s*[-_ ]\s*([A-Za-z]{2,3})\b", re.I)

_COLUMNAS_REQUERIDAS = [
    config.BMAT_COL_POSICION, config.BMAT_COL_STREAMS, config.BMAT_COL_ISRCS,
    config.BMAT_COL_DISTRIBUIDORA,
]


# --- Fuente ---------------------------------------------------------------

def identificar_archivo(path) -> tuple:
    """"WK35-CO.xlsx" -> (35, "CO"). Error claro si el nombre no sirve."""
    nombre = Path(path).name
    m = _PATRON_ARCHIVO.search(nombre)
    if not m:
        raise ValueError(
            f"No reconozco la semana y el mercado en el nombre \"{nombre}\". "
            f"Debe llamarse como los de BMAT: WK35-CO.xlsx, WK35-CAM.xlsx..."
        )
    semana, mercado = int(m.group(1)), m.group(2).upper()
    if mercado == "PN":  # por si alguien renombra Panamá como en los reportes
        mercado = "PA"
    return semana, mercado


def buscar_fuentes(carpeta, mercados=None) -> dict:
    """{mercado: ruta} de los WK de BMAT que hay en la carpeta, solo de los
    mercados que se usan (config.bmat_mercados_activos: hoy solo CO; los
    demás WK se ignoran). Si hay de varias semanas, error: mezclar semanas
    en una misma corrida no tiene sentido."""
    activos = set(mercados or config.bmat_mercados_activos())
    encontrados = {}
    for p in sorted(Path(carpeta).iterdir()):
        if p.suffix.lower() not in (".xlsx", ".xlsm") or p.name.startswith("~$"):
            continue
        try:
            semana, mercado = identificar_archivo(p)
        except ValueError:
            continue
        if mercado in activos:
            encontrados.setdefault(semana, {})[mercado] = p
    if not encontrados:
        esperados = ", ".join(f"WK<semana>-{m}.xlsx" for m in activos)
        raise ValueError(f"No encontré archivos WK de BMAT ({esperados}) en {carpeta}.")
    if len(encontrados) > 1:
        raise ValueError(
            "La carpeta tiene archivos de varias semanas "
            f"({', '.join('WK%d' % s for s in sorted(encontrados))}). "
            "Deja solo los de la semana que quieres generar."
        )
    return next(iter(encontrados.values()))


def leer_fuente(path) -> pd.DataFrame:
    """Lee un WK de BMAT (una hoja, un track por fila)."""
    df = pd.read_excel(path)
    faltan = [c for c in _COLUMNAS_REQUERIDAS if c not in df.columns]
    if faltan:
        raise ValueError(f"A {Path(path).name} le faltan las columnas: {', '.join(faltan)}.")
    df = df[pd.to_numeric(df[config.BMAT_COL_POSICION], errors="coerce").notna()].copy()
    df[config.BMAT_COL_POSICION] = df[config.BMAT_COL_POSICION].astype(int)
    df[config.BMAT_COL_STREAMS] = pd.to_numeric(df[config.BMAT_COL_STREAMS], errors="coerce").fillna(0.0)
    return df.sort_values(config.BMAT_COL_POSICION).reset_index(drop=True)


def anio_de_la_semana(semana: int, hoy: date = None) -> int:
    """El archivo no trae el año: es el actual, salvo que la semana sea
    posterior a la de hoy (ej. correr la WK52 en enero) -> el anterior."""
    hoy = hoy or date.today()
    return hoy.year if semana <= hoy.isocalendar()[1] + 1 else hoy.year - 1


def bandas_de(mercado: str, filas_fuente: int = None) -> list:
    top = config.BMAT_MERCADOS[mercado]["top"]
    if filas_fuente is not None:
        top = min(top, filas_fuente)
    return [b for b in config.BMAT_BANDAS if b <= top]


# --- Titularidad compartida -------------------------------------------------

def clave_track(artista, track) -> str:
    """Artista + track sin acentos, mayúsculas ni signos: "Peso Pluma &
    Anitta" y "Peso Pluma, Anitta" dan lo mismo (la lista original tenía
    las dos formas según el país)."""
    def limpio(v):
        texto = unicodedata.normalize("NFKD", "" if v is None or (not isinstance(v, str) and pd.isna(v)) else str(v))
        return re.sub(r"[^a-z0-9]", "", texto.encode("ascii", "ignore").decode().lower())
    return f"{limpio(artista)}|{limpio(track)}"


def cargar_titularidad(path=None) -> pd.DataFrame:
    """La lista de tracks con titularidad compartida (ver
    config.BMAT_TITULARIDAD_XLSX). Vacía si el archivo no está."""
    path = Path(path or config.BMAT_TITULARIDAD_XLSX)
    columnas = ["clave", "sello_1", "pct_1", "sello_2", "pct_2"]
    if not path.exists():
        return pd.DataFrame(columns=columnas)
    df = pd.read_excel(path)
    df = df.dropna(subset=["Track", "Artista", "Sello 1", "Sello 2"])
    validos = set(config.BMAT_LABELS)
    malos = sorted((set(df["Sello 1"]) | set(df["Sello 2"])) - validos)
    if malos:
        raise ValueError(
            f"{path.name}: sellos que no existen en el reporte: {', '.join(map(str, malos))}. "
            f"Usa uno de: {', '.join(config.BMAT_LABELS)}."
        )
    return pd.DataFrame({
        "clave": [clave_track(a, t) for a, t in zip(df["Artista"], df["Track"])],
        "sello_1": df["Sello 1"], "pct_1": df["% Sello 1"].astype(float),
        "sello_2": df["Sello 2"], "pct_2": df["% Sello 2"].astype(float),
    }).drop_duplicates("clave")


def marcar_titularidad(clasificado: pd.DataFrame, titularidad: pd.DataFrame) -> pd.DataFrame:
    """Agrega "Titularidad compartida" (ej. "The Orchard 50% / Universal
    50%") a los tracks de la lista; vacío en el resto."""
    out = clasificado.copy()
    por_clave = {r.clave: r for r in titularidad.itertuples(index=False)}
    reparto, texto = [], []
    for artista, track in zip(out.get("Artista", pd.Series(index=out.index, dtype=object)),
                              out.get("Track", pd.Series(index=out.index, dtype=object))):
        r = por_clave.get(clave_track(artista, track))
        if r is None:
            reparto.append(None)
            texto.append("")
        else:
            reparto.append(((r.sello_1, r.pct_1), (r.sello_2, r.pct_2)))
            texto.append(f"{r.sello_1} {r.pct_1:.0%} / {r.sello_2} {r.pct_2:.0%}")
    out["_reparto"] = reparto
    out["Titularidad compartida"] = texto
    return out


# --- Cálculo --------------------------------------------------------------

def calcular_bandas(clasificado: pd.DataFrame, mercado: str) -> pd.DataFrame:
    """Una fila por (banda, sello) con tracks, streams (millones) y %.

    Si la fuente pasó por `marcar_titularidad`, los streams de los tracks
    compartidos se reparten entre sus dos sellos según el % de la lista (en
    vez de ir enteros al sello con el que está clasificado). El conteo de
    tracks no se reparte: el track cuenta para su sello, como siempre."""
    filas = []
    tiene_reparto = "_reparto" in clasificado.columns
    for banda in bandas_de(mercado, len(clasificado)):
        top = clasificado[clasificado[config.BMAT_COL_POSICION] <= banda]
        total = float(top[config.BMAT_COL_STREAMS].sum())
        conteo = top.groupby("Sello").size()
        streams = top.groupby("Sello")[config.BMAT_COL_STREAMS].sum().to_dict()
        if tiene_reparto:
            for sello, s, reparto in zip(top["Sello"], top[config.BMAT_COL_STREAMS], top["_reparto"]):
                if not reparto:
                    continue
                streams[sello] = streams.get(sello, 0.0) - s
                for sello_i, pct_i in reparto:
                    streams[sello_i] = streams.get(sello_i, 0.0) + pct_i * s
        for sello in config.BMAT_LABELS:
            s = float(streams.get(sello, 0.0))
            filas.append({
                "banda": banda, "label_group": sello,
                "tracks": int(conteo.get(sello, 0)),
                "streams_millones": s / 1e6,
                "pct_streams": s / total if total else 0.0,
            })
    return pd.DataFrame(filas)


# --- Histórico ------------------------------------------------------------

def asegurar_semilla(csv_path: Path = None) -> int:
    """Siembra el histórico BMAT de los 11 mercados (semana 35 de 2026 y
    antes, sacado de los reportes reales) si todavía no está. INSERT OR
    IGNORE: nunca pisa lo ya guardado ni lo calculado después."""
    csv_path = Path(csv_path or SEED_BMAT_CSV)
    if not csv_path.exists():
        return 0
    semilla = pd.read_csv(csv_path)
    conn = history.conectar()
    try:
        hay = conn.execute(
            "SELECT COUNT(*) FROM bmat_weekly WHERE anio * 100 + semana <= ?",
            (int((semilla["anio"] * 100 + semilla["semana"]).max()),),
        ).fetchone()[0]
        if hay >= len(semilla):
            return 0
        antes = conn.total_changes
        conn.executemany(
            """INSERT OR IGNORE INTO bmat_weekly
               (anio, semana, country_code, banda, label_group,
                tracks, streams_millones, pct_streams)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
            semilla[["anio", "semana", "country_code", "banda", "label_group",
                     "tracks", "streams_millones", "pct_streams"]]
            .astype(object).where(pd.notna(semilla), None).itertuples(index=False, name=None),
        )
        conn.commit()
        return conn.total_changes - antes
    finally:
        conn.close()


def guardar_semana(anio: int, semana: int, mercado: str, bandas: pd.DataFrame) -> None:
    """Guarda (o reemplaza, si se vuelve a correr) una semana de un mercado."""
    conn = history.conectar()
    try:
        conn.execute("DELETE FROM bmat_weekly WHERE anio = ? AND semana = ? AND country_code = ?",
                     (anio, semana, mercado))
        conn.executemany(
            """INSERT INTO bmat_weekly
               (anio, semana, country_code, banda, label_group,
                tracks, streams_millones, pct_streams)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
            [(anio, semana, mercado, int(r.banda), r.label_group, int(r.tracks),
              float(r.streams_millones), float(r.pct_streams))
             for r in bandas.itertuples(index=False)],
        )
        conn.commit()
    finally:
        conn.close()


def ultima_semana(mercado: str = None):
    """(anio, semana) más reciente guardada (de un mercado o de cualquiera)."""
    conn = history.conectar()
    try:
        sql = "SELECT anio, semana FROM bmat_weekly"
        args = ()
        if mercado:
            sql += " WHERE country_code = ?"
            args = (mercado,)
        fila = conn.execute(sql + " ORDER BY anio DESC, semana DESC LIMIT 1", args).fetchone()
        return tuple(fila) if fila else None
    finally:
        conn.close()


# --- Archivo intermedio ------------------------------------------------------

_AZUL = PatternFill("solid", start_color=config.COLOR_BANNER_BMAT, end_color=config.COLOR_BANNER_BMAT)
_BLANCA = Font(bold=True, color="FFFFFF")
_NEGRITA = Font(bold=True)
_FINO = Side(style="thin", color="808080")
_BORDE = Border(left=_FINO, right=_FINO, top=_FINO, bottom=_FINO)
_CENTRO = Alignment(horizontal="center", vertical="center", wrap_text=True)

# Orden y nombres de "Streams Catalogo", como en la plantilla, con sus grupos.
_GRUPOS_CATALOGO = [
    ("Universal + Ingrooves + Virgin", ["Universal", "Virgin", "Ingrooves"]),
    ("Sony + Orchard", ["Sony", "The Orchard"]),
    ("Warner + ADA", ["Warner", "ADA Music"]),
    (None, ["Independientes"]),
]
_NOMBRE_LARGO = {v: k for k, v in config.BMAT_DISQUERAS_MAJOR.items()}
_NOMBRE_LARGO["Independientes"] = "Independientes"


def nombre_intermedio(mercado: str, anio: int, semana: int) -> str:
    m = config.BMAT_MERCADOS[mercado]
    return f"Top {m['top']} BMAT {m['sigla']} a sem {semana} de {anio}.xlsx"


def _encabezado(ws, fila, columnas):
    for j, (nombre, ancho) in enumerate(columnas, start=1):
        c = ws.cell(row=fila, column=j, value=nombre)
        c.font, c.fill, c.alignment, c.border = _BLANCA, _AZUL, _CENTRO, _BORDE
        if ancho:
            ws.column_dimensions[get_column_letter(j)].width = ancho


def _titulo(ws, texto, subtitulo=None):
    ws["A1"] = texto
    ws["A1"].font = Font(bold=True, size=13)
    if subtitulo:
        ws["A2"] = subtitulo
        ws["A2"].font = Font(italic=True, color="595959")


def _listado(ws, df, fila0, columnas_df):
    for i, fila in enumerate(df[columnas_df].itertuples(index=False), start=fila0):
        for j, v in enumerate(fila, start=1):
            c = ws.cell(row=i, column=j, value=None if (not isinstance(v, str) and pd.isna(v)) else v)
            c.border = _BORDE


def _movimiento(df):
    return df[config.BMAT_COL_MOVIMIENTO] if config.BMAT_COL_MOVIMIENTO in df else pd.Series("", index=df.index)


def _hoja_resumen(ws, bandas, etiqueta):
    _titulo(ws, f"Resumen BMAT - {etiqueta}", "Tracks, streams (millones) y % de streams por sello y banda.")
    lista_bandas = sorted(bandas["banda"].unique())
    fila = 4
    for nombre, campo, formato in [("TRACKS", "tracks", "#,##0"),
                                   ("Streams (Millones)", "streams_millones", "#,##0.00"),
                                   ("% Streams", "pct_streams", "0.0%")]:
        _encabezado(ws, fila, [(nombre, 18)] + [(f"Top {b:,}".replace(",", "."), 12) for b in lista_bandas])
        for k, sello in enumerate(config.BMAT_LABELS + ["Total"], start=1):
            c = ws.cell(row=fila + k, column=1, value=sello)
            c.border = _BORDE
            if sello == "Total":
                c.font = _NEGRITA
            for j, b in enumerate(lista_bandas, start=2):
                sub = bandas[bandas["banda"] == b]
                v = (sub[campo].sum() if sello == "Total"
                     else sub.loc[sub["label_group"] == sello, campo].sum())
                cel = ws.cell(row=fila + k, column=j, value=float(v))
                cel.number_format, cel.border = formato, _BORDE
                if sello == "Total":
                    cel.font = _NEGRITA
        fila += len(config.BMAT_LABELS) + 3
    ws.freeze_panes = "B5"
    return fila


def _nota_titularidad(ws, df, fila):
    """Debajo del resumen, los tracks de la semana con titularidad
    compartida y cómo se repartieron sus streams (lo que en la plantilla
    eran sumas a mano en las fórmulas de la hoja "Resumen")."""
    if "Titularidad compartida" not in df.columns:
        return
    comp = df[df["Titularidad compartida"] != ""]
    ws.cell(row=fila, column=1, value="Tracks con titularidad compartida (sus streams se reparten)").font = _NEGRITA
    if comp.empty:
        ws.cell(row=fila + 1, column=1, value="Ninguno en esta semana.")
        return
    _encabezado(ws, fila + 1, [("Posición", None), ("Track", None), ("Artista", None),
                               ("Streams", None), ("Clasificado como", None), ("Reparto", None)])
    for i, (_, r) in enumerate(comp.iterrows(), start=fila + 2):
        valores = [int(r[config.BMAT_COL_POSICION]), r.get("Track"), r.get("Artista"),
                   float(r[config.BMAT_COL_STREAMS]), r["Sello"], r["Titularidad compartida"]]
        for j, v in enumerate(valores, start=1):
            c = ws.cell(row=i, column=j, value=v)
            c.border = _BORDE
            if j == 4:
                c.number_format = "#,##0"


def _hoja_nuevos(ws, df):
    _titulo(ws, "TOP 200 - Tracks nuevos", "Tracks del Top 200 que entran por primera vez (\"NUEVO\").")
    mov = _movimiento(df).astype(str).str.upper()
    nuevos = df[(df[config.BMAT_COL_POSICION] <= 200) & mov.str.contains("NUEVO|NEW")].copy()
    nuevos["#"] = _movimiento(nuevos)
    _encabezado(ws, 4, [("Posición", 10), ("#", 9), ("Track", 40), ("Artista", 36),
                        ("Disqueras", 30), ("Sello", 16)])
    _listado(ws, nuevos, 5, [config.BMAT_COL_POSICION, "#", "Track", "Artista", "Disqueras", "Sello"])
    if nuevos.empty:
        ws["A5"] = "No hay tracks nuevos en el Top 200 esta semana."


def _hoja_independientes(ws, df):
    _titulo(ws, "TOP 200 - Tracks Independientes")
    ind = df[(df[config.BMAT_COL_POSICION] <= 200) & (df["Sello"] == "Independientes")].copy()
    ind["#"] = _movimiento(ind)
    _encabezado(ws, 4, [("Posición", 10), ("#", 9), ("Track", 40), ("Artista", 40), ("Disqueras", 36)])
    _listado(ws, ind, 5, [config.BMAT_COL_POSICION, "#", "Track", "Artista", "Disqueras"])


def _hoja_top50(ws, df):
    """Dos bloques, como la plantilla: los tracks de Universal Music Group
    del Top 50, y abajo el Top 50 completo con su sello."""
    top = df[df[config.BMAT_COL_POSICION] <= 50].copy()
    top["#"] = _movimiento(top)
    _titulo(ws, "TOP 50 UMG")
    umg = top[top["Disqueras"] == "Universal Music Group"]
    _encabezado(ws, 4, [("Posición", 10), ("#", 9), ("Track", 40), ("Artista", 40), ("Disqueras", 30),
                        ("Sello", 16)])
    _listado(ws, umg, 5, [config.BMAT_COL_POSICION, "#", "Track", "Artista", "Disqueras", "Sello"])
    fila = 5 + max(len(umg), 1) + 3
    ws.cell(row=fila, column=1, value="TOP 50 - POSICIONES").font = Font(bold=True, size=13)
    _encabezado(ws, fila + 2, [("Posición", None), ("#", None), ("Track", None), ("Artista", None),
                               ("Disqueras", None), ("Sello", None)])
    _listado(ws, top, fila + 3, [config.BMAT_COL_POSICION, "#", "Track", "Artista", "Disqueras", "Sello"])


def streams_catalogo(df: pd.DataFrame, corte: str = None) -> pd.DataFrame:
    """Streams de catálogo (lanzados hasta la fecha de corte) y front line
    (después) por sello, sobre toda la fuente. Los tracks sin fecha van
    aparte, para que los totales cuadren.

    Corrige un error de la plantilla: su SUMIFS tenía el rango de
    "Disqueras" relativo, y al copiarlo hacia abajo se corría una fila por
    sello -- de Virgin para abajo sumaba contra la fila equivocada (ej.
    Warner semana 35 CO: 69,9 M en la plantilla, 56,1 M bien calculado)."""
    corte = pd.Timestamp(corte or config.BMAT_CORTE_CATALOGO)
    fecha = pd.to_datetime(df.get(config.BMAT_COL_LANZAMIENTO), errors="coerce")
    tipo = pd.Series("Sin fecha", index=df.index)
    tipo[fecha <= corte] = "Catalogo"
    tipo[fecha > corte] = "FrontLine"
    tabla = (df.assign(_tipo=tipo).groupby(["Sello", "_tipo"])[config.BMAT_COL_STREAMS].sum()
             .unstack(fill_value=0.0).reindex(config.BMAT_LABELS, fill_value=0.0))
    for col in ["Catalogo", "FrontLine", "Sin fecha"]:
        if col not in tabla:
            tabla[col] = 0.0
    return tabla[["Catalogo", "FrontLine", "Sin fecha"]]


def _hoja_catalogo(ws, df, etiqueta, top):
    corte = pd.Timestamp(config.BMAT_CORTE_CATALOGO)
    tabla = streams_catalogo(df)
    _titulo(ws, f"Catálogo (lanzados hasta el {corte:%d-%m-%Y}) vs Front Line - {etiqueta}",
            f"Sobre el Top {top:,} completo.".replace(",", "."))
    _encabezado(ws, 4, [(f"TOP {top:,}".replace(",", "."), 32), ("Streams Catálogo", 18),
                        ("Streams Front Line", 18), ("Sin fecha de lanzamiento", 18)])
    orden = [s for _, sellos in _GRUPOS_CATALOGO for s in sellos]
    for i, sello in enumerate(orden, start=5):
        ws.cell(row=i, column=1, value=_NOMBRE_LARGO[sello]).border = _BORDE
        for j, col in enumerate(["Catalogo", "FrontLine", "Sin fecha"], start=2):
            c = ws.cell(row=i, column=j, value=float(tabla.loc[sello, col]))
            c.number_format, c.border = "#,##0", _BORDE
    fila = 5 + len(orden) + 1
    _encabezado(ws, fila, [("% del total", None), ("% Catálogo", None), ("% Front Line", None)])
    totales = {col: float(tabla[col].sum()) for col in ["Catalogo", "FrontLine"]}
    fila += 1
    for grupo, sellos in _GRUPOS_CATALOGO:
        filas_grupo = ([(grupo, sellos, True)] if grupo else []) + [(_NOMBRE_LARGO[s], [s], False) for s in sellos]
        for nombre, miembros, es_grupo in filas_grupo:
            ws.cell(row=fila, column=1, value=nombre).border = _BORDE
            if es_grupo:
                ws.cell(row=fila, column=1).font = _NEGRITA
            for j, col in enumerate(["Catalogo", "FrontLine"], start=2):
                v = float(tabla.loc[miembros, col].sum()) / totales[col] if totales[col] else 0.0
                c = ws.cell(row=fila, column=j, value=v)
                c.number_format, c.border = "0.0%", _BORDE
                if es_grupo:
                    c.font = _NEGRITA
            fila += 1
    ws.cell(row=fila, column=1, value="Total").font = _NEGRITA
    for j in (2, 3):
        c = ws.cell(row=fila, column=j, value=1.0 if any(totales.values()) else 0.0)
        c.number_format, c.font = "0.0%", _NEGRITA


def _hoja_base(ws, df):
    extra = [c for c in ("Disqueras", "Sello", "Clasificacion", "Titularidad compartida") if c in df.columns]
    columnas = [c for c in df.columns
                if not str(c).startswith(("Unnamed", "_")) and c not in extra] + extra
    _encabezado(ws, 1, [(c, None) for c in columnas])
    for i, fila in enumerate(df[columnas].itertuples(index=False), start=2):
        for j, v in enumerate(fila, start=1):
            if not isinstance(v, str) and pd.isna(v):
                continue
            ws.cell(row=i, column=j, value=v.item() if hasattr(v, "item") else v)
    anchos = {"Track": 36, "Artista": 32, "ISRCs": 28, "Disquera original": 28,
              "Distribuidora original": 26, "Disqueras": 26}
    for j, c in enumerate(columnas, start=1):
        ws.column_dimensions[get_column_letter(j)].width = anchos.get(c, 12)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columnas))}{len(df) + 1}"


def escribir_intermedio(path: Path, mercado: str, anio: int, semana: int,
                        clasificado: pd.DataFrame, bandas: pd.DataFrame,
                        hojas_a_generar=None) -> Path:
    """El archivo "Top N BMAT XX a sem S de A": lo que antes se armaba a
    mano con tablas dinámicas, ya calculado. Lleva las hojas de
    config.BMAT_HOJAS_INTERMEDIO, en ese orden (o las que se pidan)."""
    m = config.BMAT_MERCADOS[mercado]
    etiqueta = f"{m['nombre']} (sem {semana} de {anio})"
    familia_b = mercado in config.BMAT_REPORTES["CAM"]["hojas"]
    top = min(m["top"], len(clasificado))

    def resumen(ws):
        _nota_titularidad(ws, clasificado, _hoja_resumen(ws, bandas, etiqueta))

    hojas = {
        "Streams Catalogo": lambda ws: _hoja_catalogo(ws, clasificado, etiqueta, top),
        "Resumen": resumen,
        "TOP 200 Nuevos": lambda ws: _hoja_nuevos(ws, clasificado),
        "Tracks Independientes": lambda ws: _hoja_independientes(ws, clasificado),
        "TOP 50 - Posiciones UMG": lambda ws: _hoja_top50(ws, clasificado),
        "Archivo Base": lambda ws: _hoja_base(ws, clasificado),
    }
    wb = Workbook()
    wb.remove(wb.active)
    for nombre in (config.BMAT_HOJAS_INTERMEDIO if hojas_a_generar is None else hojas_a_generar):
        titulo_hoja = "Streams Cat&Front" if (nombre == "Streams Catalogo" and familia_b) else nombre
        hojas[nombre](wb.create_sheet(titulo_hoja))
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
