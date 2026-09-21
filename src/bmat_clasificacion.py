"""Clasificación de los tracks BMAT: a qué sello va cada canción.

Es el único paso del proceso BMAT que no es mecánico. En el proceso manual,
al pegar la fuente en "Archivo Base" se le agregaba la columna "Disqueras":
el nombre de la major dueña del track ("Universal Music Group", "Virgin",
"INgrooves"...) o, si es independiente, el de su distribuidora ("oneRPM",
"DistroKid"...). Con esa columna salen los 8 sellos del reporte (ver
config.BMAT_DISQUERAS_MAJOR).

Lo que se aprendió de los intermedios de la semana 35 (11 mercados, 25.122
ISRC):

- La clasificación es por track y es la MISMA en todos los países: ningún
  ISRC tiene dos "Disqueras" distintas.
- ~92% de los tracks van a su "Distribuidora original", pero el resto se
  reasigna a mano, track por track (no hay una regla fija).
- ~96% de los tracks se repiten de una semana a la otra.

Por eso el esquema es una TABLA por ISRC, compartida por todos los mercados:

1. Se siembra una vez con la semana 35 (`seed_bmat_clasificacion.csv.gz`).
2. Cada semana, un track cuyo ISRC ya está en la tabla toma su clasificación
   de ahí. Uno nuevo toma la de la regla por defecto -- la "Disqueras" más
   común entre los tracks conocidos de su misma distribuidora, que acierta
   ~97% -- y sale en la "Lista para revisar".
3. Lo que el usuario corrija en esa lista se guarda en la tabla
   (`aplicar_revision`) y vale para todas las semanas siguientes.

Un track trae a veces varios ISRC ("NL8UF2614553, QZNJW2601079"): se guarda
cada uno por separado, y al clasificar basta con que uno sea conocido.
"""
from pathlib import Path
from datetime import datetime

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import config, history

SEED_CLASIFICACION_CSV = history.SEED_DIR / "seed_bmat_clasificacion.csv.gz"

# Origen de cada clasificación guardada.
ORIGEN_SEMILLA = "semilla"      # tomada de los intermedios de la semana 35
ORIGEN_REGLA = "regla"          # asignada por la regla por defecto, sin revisar
ORIGEN_REVISION = "revision"    # corregida o confirmada por el usuario

HOJA_REVISION = "Lista para revisar"
COLUMNA_CORRECCION = "Disqueras corregida"

_MAJOR_POR_MINUSCULA = {k.lower(): k for k in config.BMAT_DISQUERAS_MAJOR}


def _crear_tabla(conn) -> None:
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS bmat_clasificacion (
            isrc TEXT PRIMARY KEY,
            disqueras TEXT NOT NULL,
            origen TEXT NOT NULL,
            track TEXT,
            artista TEXT,
            disquera_original TEXT,
            distribuidora_original TEXT,
            actualizado TEXT
        )
        """
    )


def normalizar_disqueras(valor) -> str:
    """Quita espacios y unifica la grafía de las majors ("Ingrooves" ->
    "INgrooves"); cualquier otro nombre queda tal cual."""
    texto = str(valor).strip() if valor is not None and not pd.isna(valor) else ""
    return _MAJOR_POR_MINUSCULA.get(texto.lower(), texto)


def sello_de(disqueras) -> str:
    """"Disqueras" -> uno de los 8 sellos del reporte."""
    return config.BMAT_DISQUERAS_MAJOR.get(normalizar_disqueras(disqueras), "Independientes")


def separar_isrcs(valor) -> list:
    """"NL8UF2614553, QZNJW2601079" -> ["NL8UF2614553", "QZNJW2601079"]."""
    if valor is None or (not isinstance(valor, str) and pd.isna(valor)):
        return []
    return [p.strip().upper() for p in str(valor).replace(";", ",").split(",") if p.strip()]


def sembrar(csv_path: Path = None) -> int:
    """Carga la clasificación de la semana 35 si la tabla está vacía.
    Nunca pisa lo que ya haya (INSERT OR IGNORE): las correcciones del
    usuario valen más que la semilla. Devuelve cuántas filas agregó."""
    csv_path = csv_path or SEED_CLASIFICACION_CSV
    conn = history.conectar()
    try:
        _crear_tabla(conn)
        if conn.execute("SELECT COUNT(*) FROM bmat_clasificacion").fetchone()[0]:
            return 0
        if not Path(csv_path).exists():
            return 0
        df = pd.read_csv(csv_path, dtype=str).fillna("")
        antes = conn.total_changes
        conn.executemany(
            """INSERT OR IGNORE INTO bmat_clasificacion
               (isrc, disqueras, origen, track, artista, disquera_original,
                distribuidora_original, actualizado)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
            [
                (r.isrc.strip().upper(), normalizar_disqueras(r.disqueras), ORIGEN_SEMILLA,
                 r.track, r.artista, r.disquera_original, r.distribuidora_original, "semana 35 de 2026")
                for r in df.itertuples(index=False)
            ],
        )
        conn.commit()
        return conn.total_changes - antes
    finally:
        conn.close()


def cargar() -> pd.DataFrame:
    conn = history.conectar()
    try:
        _crear_tabla(conn)
        return pd.read_sql_query("SELECT * FROM bmat_clasificacion", conn)
    finally:
        conn.close()


def regla_por_distribuidora(tabla: pd.DataFrame) -> dict:
    """{distribuidora original: ("Disqueras" más común entre sus tracks
    conocidos, qué fracción de ellos la tiene)}. En empate gana la primera
    en orden alfabético, para que el resultado no dependa del orden de la
    tabla."""
    if tabla.empty:
        return {}
    conteo = (tabla[tabla["distribuidora_original"].fillna("") != ""]
              .groupby(["distribuidora_original", "disqueras"]).size()
              .reset_index(name="n"))
    conteo["share"] = conteo["n"] / conteo.groupby("distribuidora_original")["n"].transform("sum")
    conteo = conteo.sort_values(["distribuidora_original", "n", "disqueras"], ascending=[True, False, True])
    mejor = conteo.drop_duplicates("distribuidora_original")
    return {d: (q, float(s)) for d, q, s in zip(mejor["distribuidora_original"], mejor["disqueras"], mejor["share"])}


def disqueras_por_defecto(distribuidora, regla: dict) -> tuple:
    """(clasificación, confianza) de un track nuevo: la regla de su
    distribuidora, con la fracción de tracks conocidos de esa distribuidora
    que la tienen; si la distribuidora nunca se vio, ella misma, sin
    confianza (None)."""
    texto = normalizar_disqueras(distribuidora)
    if texto in regla:
        disqueras, share = regla[texto]
        return normalizar_disqueras(disqueras), share
    return (texto or "Sin distribuidora"), None


def clasificar(df: pd.DataFrame, tabla: pd.DataFrame = None, etiqueta: str = None) -> pd.DataFrame:
    """Agrega a una fuente WK las columnas "Disqueras", "Sello" y
    "Clasificacion". No guarda nada.

    "Clasificacion" es "conocido" (su ISRC ya estaba clasificado: semilla,
    revisión del usuario o regla de una semana anterior) o "nuevo" (se le
    aplicó la regla por defecto ahora). Un track que entró por la regla en
    ESTA misma semana (`etiqueta`) sigue contando como nuevo si se vuelve a
    correr la semana: así la lista para revisar sale igual la segunda vez,
    menos lo que el usuario ya corrigió."""
    tabla = cargar() if tabla is None else tabla
    vigentes = tabla
    if etiqueta is not None and not tabla.empty:
        vigentes = tabla[~((tabla["origen"] == ORIGEN_REGLA) & (tabla["actualizado"] == etiqueta))]
    por_isrc = dict(zip(vigentes["isrc"], vigentes["disqueras"]))
    regla = regla_por_distribuidora(vigentes)

    disqueras, estado, confianza = [], [], []
    for isrcs, distribuidora in zip(df[config.BMAT_COL_ISRCS], df[config.BMAT_COL_DISTRIBUIDORA]):
        conocidas = [por_isrc[i] for i in separar_isrcs(isrcs) if i in por_isrc]
        if conocidas:
            # Si los ISRC de un mismo track discreparan, gana el más repetido.
            disqueras.append(normalizar_disqueras(pd.Series(conocidas).mode().iloc[0]))
            estado.append("conocido")
            confianza.append(None)
        else:
            d, share = disqueras_por_defecto(distribuidora, regla)
            disqueras.append(d)
            estado.append("nuevo")
            confianza.append(share)
    out = df.copy()
    out["Disqueras"] = disqueras
    out["Sello"] = [sello_de(d) for d in disqueras]
    out["Clasificacion"] = estado
    out["_confianza"] = confianza
    return out


def guardar_nuevos(clasificado: pd.DataFrame, etiqueta: str) -> int:
    """Guarda en la tabla, con origen "regla" y la etiqueta de la semana,
    los tracks nuevos, para que la semana siguiente ya sean conocidos (y no
    vuelvan a salir en la lista). Nunca pisa una clasificación de la
    semilla ni una revisada por el usuario."""
    nuevos = clasificado[clasificado["Clasificacion"] == "nuevo"]
    filas = []
    for _, r in nuevos.iterrows():
        for isrc in separar_isrcs(r[config.BMAT_COL_ISRCS]):
            filas.append((isrc, r["Disqueras"], ORIGEN_REGLA, _texto(r.get("Track")),
                          _texto(r.get("Artista")), _texto(r.get(config.BMAT_COL_DISQUERA)),
                          _texto(r.get(config.BMAT_COL_DISTRIBUIDORA)), etiqueta))
    conn = history.conectar()
    try:
        _crear_tabla(conn)
        antes = conn.total_changes
        conn.executemany(
            """INSERT INTO bmat_clasificacion
               (isrc, disqueras, origen, track, artista, disquera_original,
                distribuidora_original, actualizado)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?)
               ON CONFLICT(isrc) DO UPDATE SET
                 disqueras = excluded.disqueras, actualizado = excluded.actualizado
               WHERE bmat_clasificacion.origen = 'regla'
                 AND bmat_clasificacion.actualizado = excluded.actualizado""",
            filas,
        )
        conn.commit()
        return conn.total_changes - antes
    finally:
        conn.close()


def _texto(v) -> str:
    return "" if v is None or (not isinstance(v, str) and pd.isna(v)) else str(v)


# --- Lista para revisar ------------------------------------------------------

# Un track nuevo se marca "Alta" prioridad si está en el Top 200 de algún
# mercado (mueve las bandas que más se miran) o si la regla de su
# distribuidora es poco confiable. Con la semana 35 (simulando tracks
# nuevos) la regla acertó 97%, pero con confianza menor a 90% falla 4 de
# cada 10: ahí se concentra la mitad de los errores en ~4% de los tracks.
PRIORIDAD_TOP = 200
PRIORIDAD_CONFIANZA = 0.9

_COLUMNAS_REVISION = [
    ("Prioridad", 10), ("Mejor posición", 11), ("Mercados", 18), ("Streams (total)", 14), ("Track", 38),
    ("Artista", 32), ("ISRCs", 30), ("Disquera original", 30), ("Distribuidora original", 28),
    ("Disqueras asignada", 26), ("Confianza regla", 11), ("Sello asignado", 15), (COLUMNA_CORRECCION, 26),
]


def lista_para_revisar(clasificados: dict) -> pd.DataFrame:
    """Une los tracks nuevos de todos los mercados de la semana (uno por
    fila aunque salga en varios países), ordenados por su mejor posición:
    los de arriba son los que más mueven el market share."""
    partes = []
    for mercado, df in clasificados.items():
        nuevos = df[df["Clasificacion"] == "nuevo"].copy()
        if nuevos.empty:
            continue
        nuevos["_mercado"] = config.BMAT_MERCADOS[mercado]["sigla"]
        partes.append(nuevos)
    if not partes:
        return pd.DataFrame(columns=[c for c, _ in _COLUMNAS_REVISION])
    todo = pd.concat(partes, ignore_index=True)
    todo["_clave"] = todo[config.BMAT_COL_ISRCS].map(lambda v: ", ".join(sorted(separar_isrcs(v))))
    grupos = []
    for _, g in todo.groupby("_clave", sort=False):
        primera = g.sort_values(config.BMAT_COL_POSICION).iloc[0]
        mejor = int(g[config.BMAT_COL_POSICION].min())
        confianza = primera.get("_confianza")
        confianza = None if confianza is None or pd.isna(confianza) else float(confianza)
        alta = mejor <= PRIORIDAD_TOP or confianza is None or confianza < PRIORIDAD_CONFIANZA
        grupos.append({
            "Prioridad": "Alta" if alta else "Normal",
            "Mejor posición": mejor,
            "Mercados": ", ".join(dict.fromkeys(g["_mercado"])),
            "Streams (total)": float(pd.to_numeric(g[config.BMAT_COL_STREAMS], errors="coerce").sum()),
            "Track": _texto(primera.get("Track")),
            "Artista": _texto(primera.get("Artista")),
            "ISRCs": _texto(primera[config.BMAT_COL_ISRCS]),
            "Disquera original": _texto(primera.get(config.BMAT_COL_DISQUERA)),
            "Distribuidora original": _texto(primera.get(config.BMAT_COL_DISTRIBUIDORA)),
            "Disqueras asignada": primera["Disqueras"],
            "Confianza regla": confianza,
            "Sello asignado": primera["Sello"],
            COLUMNA_CORRECCION: None,
        })
    return pd.DataFrame(grupos).sort_values(["Prioridad", "Mejor posición", "Streams (total)"],
                                            ascending=[True, True, False]).reset_index(drop=True)


def escribir_lista_revision(lista: pd.DataFrame, path: Path, titulo: str) -> Path:
    """El Excel que el usuario corrige. Solo hay que llenar la última
    columna en los tracks mal asignados (con el nombre de una major, que se
    elige de la lista desplegable, o el de la distribuidora si es
    independiente) y volver a generar con ese archivo."""
    wb = Workbook()
    ws = wb.active
    ws.title = HOJA_REVISION
    ws["A1"] = titulo
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = (
        f"Tracks que no estaban clasificados: se asignaron con la regla por distribuidora "
        f"(\"Confianza regla\" = qué parte de los tracks conocidos de esa distribuidora va a esa "
        f"clasificación). Empieza por los de prioridad Alta. "
        f"Si alguno está mal, escribe la clasificación correcta en \"{COLUMNA_CORRECCION}\" "
        f"(una major de la lista, o el nombre de la distribuidora si es independiente) "
        f"y vuelve a generar la semana eligiendo este archivo. Lo que dejes vacío queda como está."
    )
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(_COLUMNAS_REVISION))
    ws.row_dimensions[2].height = 45

    azul = PatternFill("solid", start_color=config.COLOR_BANNER_BMAT, end_color=config.COLOR_BANNER_BMAT)
    amarillo = PatternFill("solid", start_color="FFF2CC", end_color="FFF2CC")
    fila_enc = 4
    for j, (nombre, ancho) in enumerate(_COLUMNAS_REVISION, start=1):
        c = ws.cell(row=fila_enc, column=j, value=nombre)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = azul
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = ancho
    for i, fila in enumerate(lista.itertuples(index=False), start=fila_enc + 1):
        for j, valor in enumerate(fila, start=1):
            c = ws.cell(row=i, column=j, value=None if pd.isna(valor) else valor)
            nombre = _COLUMNAS_REVISION[j - 1][0]
            if nombre == "Streams (total)":
                c.number_format = "#,##0"
            elif nombre == "Confianza regla":
                c.number_format = "0%"
            elif nombre == "Prioridad" and valor == "Alta":
                c.font = Font(bold=True, color="C00000")
        ws.cell(row=i, column=len(_COLUMNAS_REVISION)).fill = amarillo

    ultima = fila_enc + max(len(lista), 1)
    col_corr = get_column_letter(len(_COLUMNAS_REVISION))
    # Lista desplegable con las majors, pero sin bloquear otro texto (un
    # independiente se escribe con el nombre de su distribuidora).
    dv = DataValidation(type="list", formula1='"' + ",".join(config.BMAT_DISQUERAS_MAJOR) + '"',
                        allow_blank=True, showErrorMessage=False)
    dv.add(f"{col_corr}{fila_enc + 1}:{col_corr}{ultima + 200}")
    ws.add_data_validation(dv)
    ws.freeze_panes = ws.cell(row=fila_enc + 1, column=4).coordinate
    ws.auto_filter.ref = f"A{fila_enc}:{col_corr}{ultima}"
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def aplicar_revision(path: Path) -> int:
    """Lee una "Lista para revisar" corregida y guarda cada corrección en la
    tabla (origen "revision"), para todos los ISRC del track. Devuelve
    cuántos tracks se corrigieron."""
    wb = load_workbook(path, read_only=True, data_only=True)
    if HOJA_REVISION not in wb.sheetnames:
        raise ValueError(f"El archivo {Path(path).name} no tiene la hoja \"{HOJA_REVISION}\".")
    filas = list(wb[HOJA_REVISION].iter_rows(values_only=True))
    wb.close()
    encabezado = next((i for i, f in enumerate(filas) if f and COLUMNA_CORRECCION in f), None)
    if encabezado is None:
        raise ValueError(f"No encontré la columna \"{COLUMNA_CORRECCION}\" en {Path(path).name}.")
    nombres = list(filas[encabezado])
    i_corr, i_isrc = nombres.index(COLUMNA_CORRECCION), nombres.index("ISRCs")
    i_track, i_art = nombres.index("Track"), nombres.index("Artista")
    i_disq, i_dist = nombres.index("Disquera original"), nombres.index("Distribuidora original")
    ahora = datetime.now().strftime("%Y-%m-%d %H:%M")

    registros, tracks = [], 0
    for f in filas[encabezado + 1:]:
        if not f or len(f) <= i_corr:
            continue
        correccion = normalizar_disqueras(f[i_corr])
        isrcs = separar_isrcs(f[i_isrc])
        if not correccion or not isrcs:
            continue
        tracks += 1
        for isrc in isrcs:
            registros.append((isrc, correccion, ORIGEN_REVISION, _texto(f[i_track]), _texto(f[i_art]),
                              _texto(f[i_disq]), _texto(f[i_dist]), ahora))
    conn = history.conectar()
    try:
        _crear_tabla(conn)
        conn.executemany(
            """INSERT INTO bmat_clasificacion
               (isrc, disqueras, origen, track, artista, disquera_original,
                distribuidora_original, actualizado)
               VALUES (?, ?, ?, ?, ?, ?, ?, ?)
               ON CONFLICT(isrc) DO UPDATE SET
                 disqueras = excluded.disqueras, origen = excluded.origen,
                 actualizado = excluded.actualizado""",
            registros,
        )
        conn.commit()
    finally:
        conn.close()
    return tracks
